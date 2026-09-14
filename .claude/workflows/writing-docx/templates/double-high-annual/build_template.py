"""Administrator build of the current-project annual report template."""
from copy import deepcopy
from hashlib import sha256
import json
from pathlib import Path
import sys

from docx import Document
from docx.shared import Pt, Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

DIRECTORY = Path(__file__).resolve().parent
ROOT = DIRECTORY.parents[4]
sys.path.insert(0, str(ROOT / "python"))
from reporting.bindings import document_xml, text_hash, NS


def build():
    reference = next(p for p in (DIRECTORY.parent / "szpt-midterm").glob("*规范化模板.docx") if not p.name.startswith("~$"))
    document = Document(reference)
    body = document._element.body
    for element in list(body):
        if element.tag != qn("w:sectPr"):
            body.remove(element)
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(2.2), Cm(2.8)
    section.left_margin = section.right_margin = Cm(2.2)
    for style_name, size in [("Normal",11),("Title",22),("Heading 1",17),("Heading 2",13)]:
        style = document.styles[style_name] if style_name in document.styles else document.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
        style.font.name, style.font.size = "宋体", Pt(size)
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.line_spacing = 1.25
    slots, datasets = [], {}
    def paragraph(text, style=None): return document.add_paragraph(text, style)
    def slot(name, section_key, kind="scalar", **binding):
        node = paragraph("待填 " + name)
        value = {"slot_key":name,"section_key":section_key,"name":name,"kind":kind,"required":True,**binding}
        slots.append((node._p,value))
        return value
    def dataset(role, query, **params):
        key=role+"_"+query
        definitions={k:{"from":"literal","value":v} for k,v in params.items()}
        if query not in {"project_catalog","performance_tree","report_cycle_budget"}:
            definitions["year"]={"from":"report_parameter","key":"years","expand":"each"}
        datasets[key]={"dataset_key":key,"source_role":"hpm","domain":"hpm","scope_role":role,
                       "query_id":query,"parameter_bindings":definitions,"section_keys":[role]}
        return key
    def table(name, role, query, columns, selector=None, widths=None, **params):
        key=dataset(role,query,**params)
        value=document.add_table(rows=2,cols=len(columns))
        if "Table Grid" in document.styles: value.style="Table Grid"
        value.autofit=False
        widths = widths or [16.6/len(columns)]*len(columns)
        for column,width in zip(value.columns,widths): column.width=Cm(width)
        for cell,(field,title) in zip(value.rows[0].cells,columns):
            cell.text=title
            for run in cell.paragraphs[0].runs: run.bold=True
        repeat=OxmlElement("w:tblHeader")
        value.rows[0]._tr.get_or_add_trPr().append(repeat)
        for cell in value.rows[1].cells: cell.text="待填"
        for row in value.rows:
            for cell,width in zip(row.cells,widths):
                cell.width=Cm(width)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after=Pt(3)
                    p.paragraph_format.line_spacing=1.1
                    for run in p.runs: run.font.size=Pt(9)
        slots.append((value.rows[0].cells[0].paragraphs[0]._p,{
            "slot_key":name,"section_key":role,"name":name,"kind":"table","required":True,
            "dataset_key":key,"columns":[{"field":f,"title":t,"null_text":"不适用" if f=='current_average_progress' else "未填报"} for f,t in columns],
            "selector":selector or {},"empty_text":"无符合条件的记录"}))
    paragraph("双高项目年度建设报告", "Title")
    slot("报告截止日","overview",report_parameter="as_of",prefix="报告截止日：")
    paragraph("一、报告范围与评价依据","Heading 1")
    paragraph("本报告按所选两个双高项目及年度阶段组织建设任务、资金、绩效与反馈。专业群名称取自项目登记；建设章节取自该项目当年一级任务。项目之间分别列示，不将两个专业群之和视为学校全部建设项目。")
    paragraph("历史实绩只采用所选年度、截止日前的有效反馈。任务和绩效主表及资金表展示当前数据库登记值，供核验使用；没有历史版本依据时，不将现值追认为截止日实绩。空记录、未填报和数值零分别表达。")
    paragraph("二、年度建设情况","Heading 1")
    paragraph("各项目的建设指标、下属任务和数据覆盖情况见后续项目章节。任务平均进度按三级任务现值计算；期内反馈记录数用于说明材料覆盖，不能视为成果数量或自评得分。")
    paragraph("三、评价与后续工作","Heading 1")
    paragraph("对缺少期内反馈、历史绩效或材料依据的事项，暂不作达标、获奖、建成及质量等级判断。后续应按项目、年度和建设指标补齐审核后的事实材料，核对目标与完成值的适用期间，再开展正式绩效评价。")
    paragraph("评分、权重和自评等级本次留空。")
    for role, numeral in [("group_a","一"),("group_b","二")]:
        document.add_page_break()
        project=dataset(role,"project_catalog",name_pattern=None)
        slot(role+"_name",role,dataset_key=project,field="project_name",prefix="项目"+numeral+"：")
        paragraph("1 建设任务与反馈覆盖","Heading 2")
        paragraph("下表章节来自该年度一级任务。现值平均进度不代表历史截止日完成度；无下属三级任务时不计算平均进度。")
        table(role+"_sections",role,"report_task_sections",[
            ("section_name","一级建设指标"),("task_count","三级任务数"),
            ("current_average_progress","当前平均进度 %"),("period_feedback_count","期内反馈条数")],widths=[7,2.3,4,3.3])
        paragraph("2 资金情况","Heading 2")
        funds=dataset(role,"fund_totals")
        for field,title in [("budget_amount","年度预算"),("execute_amount","年度执行金额"),("arrival_amount","年度到位金额")]:
            slot(role+"_"+field,role,dataset_key=funds,field=field,prefix=title+"（万元）：",null_text="未填报",format={"decimal_places":2})
        cycle=dataset(role,"report_cycle_budget")
        slot(role+"_cycle",role,dataset_key=cycle,field="cycle_budget_amount",prefix="全建设期预算登记值（万元）：",null_text="未填报",format={"decimal_places":2})
        paragraph("年度资金仅统计该年度一级资金行，全建设期预算仅统计无阶段一级资金行。预算为零时不计算预算执行率；上述资金现值尚不能证明截止日历史支出。")
        document.add_page_break()
        paragraph("3 任务明细与实绩依据","Heading 2")
        table(role+"_tasks",role,"task_tree",[("task_code","编码"),("task_name","三级任务"),
            ("target_value","当前目标"),("complete_value","当前完成登记"),("current_progress","当前进度 %")],selector={"task_level":3},widths=[1.8,6.5,2.7,2.8,2.8])
        paragraph("4 历史反馈与佐证材料","Heading 2")
        table(role+"_feedback",role,"task_feedback_candidates",[("task_name","任务"),("feedback_date","业务反馈日期"),("content","期内已审核反馈")])
        paragraph("无符合条件的记录表示未取得该期间可用反馈，不表示建设成果为零。附件须依据反馈业务引用另行核实访问权限和有效性。")
        paragraph("5 绩效关联与评分","Heading 2")
        table(role+"_performance",role,"report_task_performance_links",[("task_name","关联任务"),("first_performance_name","一级绩效分类"),("second_performance_name","二级绩效分类"),("performance_name","具体绩效指标")])
        paragraph("绩效关联使用任务对应年度及项目模块阶段规则。关联为空时不借用其他年度或其他项目的绩效；历史目标、完成值及完成率仍需相应期间反馈。")
        slot(role+"_score",role,intentional_blank=True,required=False,prefix="自评得分：")
    path=DIRECTORY/"template.docx"
    document.save(path)
    xml=document_xml(path)
    # The new template has no duplicate physical paragraph identities.
    original=list(document._element.xpath(".//w:p"))
    rendered=xml.xpath(".//w:p",namespaces=NS)
    mapped=[]
    for node,value in slots:
        target=rendered[original.index(node)]
        value["locator"]={"part":"word/document.xml","path":xml.getroottree().getpath(target),"expected_text_hash":text_hash(target)}
        mapped.append(value)
    def save(name,value): (DIRECTORY/name).write_text(json.dumps(value,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    parameters=json.loads((DIRECTORY.parent/"szpt-midterm/template.json").read_text(encoding="utf-8"))["parameters"]
    parameters["properties"]["years"]["maxItems"]=1
    parameters["properties"]["period_mode"]={"const":"annual"}
    parameters["required"].append("as_of")
    save("template.json",{"template_key":"double-high-annual","name":"双高项目年度建设报告","version":1,
        "enabled":True,"capabilities":["document-writing"],"docx_file":"template.docx","docx_sha256":sha256(path.read_bytes()).hexdigest(),
        "bindings_file":"query-bindings.json","source_roles":{"hpm":"schoolDoubleHigh"},"scope_roles":{"hpm":["group_a","group_b"]},
        "missing_policy":"reject","parameters":parameters})
    save("query-bindings.json",{"template_version":1,"slot_files":["slots.json"],"datasets":list(datasets.values())})
    save("slots.json",{"slots":mapped})


if __name__ == "__main__":
    build()
