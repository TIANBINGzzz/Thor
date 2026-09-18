"""复测实际文档渲染：记录耗时、结构保留及指定页截图，不承担报告撰写。"""

import argparse
import hashlib
import json
import re
from pathlib import Path
import sys
import time
import zipfile
from xml.etree import ElementTree as ET


def structure(path):
    with zipfile.ZipFile(path) as archive:
        root = ET.fromstring(archive.read("word/document.xml"))
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        return {
            "paragraphs": len(root.findall(".//w:p", ns)),
            "tables": len(root.findall(".//w:tbl", ns)),
            "sections": len(root.findall(".//w:sectPr", ns)),
            "text": "".join(n.text or "" for n in root.findall(".//w:t", ns)),
            "media": sorted(hashlib.sha256(archive.read(n)).hexdigest()
                            for n in archive.namelist() if n.startswith("word/media/")),
        }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--pages", default="2,4,5", help="One-based PDF pages for visual review")
    parser.add_argument("--fixture", action="store_true", help="Create a new Chinese TOC fixture at source (must not exist)")
    args = parser.parse_args()
    source, output = args.source.resolve(), args.output.resolve()
    output.mkdir(parents=True, exist_ok=False)
    if args.fixture:
        if source.exists():
            raise FileExistsError(source)
        from docx import Document
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        document = Document()
        document.add_paragraph("目录")
        paragraph = document.add_paragraph()
        for field_type in ("begin", "separate", "end"):
            field = OxmlElement("w:fldChar")
            field.set(qn("w:fldCharType"), field_type)
            paragraph.add_run()._r.append(field)
            if field_type == "begin":
                instruction = OxmlElement("w:instrText")
                instruction.text = ' TOC \\o "1-3" \\h '
                paragraph.add_run()._r.append(instruction)
            elif field_type == "separate":
                paragraph.add_run("OUTDATED 999")
        document.add_page_break()
        document.add_heading("第一章 建设进展", 1)
        document.add_paragraph("中文字体验证：重庆建筑、轨道交通、阶段成果。")
        table = document.add_table(rows=2, cols=3)
        table.cell(0, 0).merge(table.cell(0, 2)).text = "合并单元格"
        table.cell(1, 0).text = "指标"
        table.cell(1, 1).text = "数量"
        table.cell(1, 2).text = "待核实"
        document.add_page_break()
        document.add_paragraph("附加验证页")
        document.add_page_break()
        document.add_heading("第二章 后续安排", 1)
        document.add_paragraph("用于验证目录页码重新计算。")
        source.parent.mkdir(parents=True, exist_ok=True)
        document.save(source)
    original = hashlib.sha256(source.read_bytes()).hexdigest()
    before = structure(source)
    start = time.monotonic()
    sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "python"))
    from tools.document_conversion import render_document
    result = render_document(str(source), str(output / "render"), str(source.parent), [str(output)])
    docx, pdf = Path(result["docxPath"]), Path(result["pdfPath"])
    result["seconds"] = round(time.monotonic() - start, 3)
    after = structure(docx)
    result.update({"sourceUnchanged": hashlib.sha256(source.read_bytes()).hexdigest() == original,
                   "before": {k: v for k, v in before.items() if k != "text"},
                   "after": {k: v for k, v in after.items() if k != "text"},
                   "mediaContentPreserved": set(before["media"]) == set(after["media"]),
                   "textContentPreserved": "".join(before["text"].split()) == "".join(after["text"].split()),
                   "sourceTextLength": len(before["text"]), "outputTextLength": len(after["text"])})
    import pypdfium2 as pdfium
    document = pdfium.PdfDocument(pdf)
    try:
        result["pageCount"] = len(document)
        result["pages"] = []
        for number in sorted({1, *(int(p) for p in args.pages.split(","))}):
            if number > len(document):
                continue
            page = document[number - 1]
            textpage = page.get_textpage()
            result["pages"].append({"page": number, "text": textpage.get_text_bounded()})
            textpage.close()
            bitmap = page.render(scale=1.3)
            image = bitmap.to_pil()
            image.save(output / f"page-{number:04}.png")
            image.close()
            bitmap.close()
            page.close()
    finally:
        document.close()
    if args.fixture:
        toc = result["pages"][0]["text"]
        result["tocVerified"] = bool(result["pageCount"] == 4 and "999" not in toc
                                     and re.search(r"第一章.*?2", toc, re.S)
                                     and re.search(r"第二章.*?4", toc, re.S))
    (output / "result.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({k: v for k, v in result.items() if k not in ("pages", "before", "after")}, ensure_ascii=False))
    if args.fixture and not result["tocVerified"]:
        raise RuntimeError("目录未更新为实际的第2页和第4页")


if __name__ == "__main__":
    main()
