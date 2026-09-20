import tempfile
import unittest
import json
import asyncio
from pathlib import Path
from unittest.mock import patch
from docx import Document

from tools.artifacts import publish_artifact, create_artifact_server


class ArtifactToolTests(unittest.TestCase):
    def test_rejects_markdown_disguised_as_word_without_publishing(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            work, output = root / 'work', root / 'output'
            work.mkdir()
            (work / 'draft.md').write_text('# 天气研究\n这不是Word文件', encoding='utf-8')
            for name in ('天气研究.docx', '天气研究.docx.md'):
                with self.subTest(name=name), self.assertRaisesRegex(ValueError, 'DOCX'):
                    publish_artifact('draft.md', name, work, output, root)
            self.assertFalse(list(output.glob('artifact_*')))

    def test_tool_receipt_has_no_upload_state_and_preserves_real_docx(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            work, output = root / 'work', root / 'output'
            work.mkdir()
            doc = Document()
            doc.add_paragraph('天气研究正文')
            doc.save(work / 'draft.docx')
            events = []
            with patch('tools.artifacts.create_sdk_mcp_server', side_effect=lambda *args, **kwargs: kwargs):
                server = create_artifact_server(root, work, output, on_published=events.append)
            (work / 'bad.docx').write_text('Markdown is not Word', encoding='utf-8')
            rejected = asyncio.run(server['tools'][0].handler({'source_path': 'bad.docx', 'file_name': '天气研究.docx'}))
            self.assertTrue(rejected['isError'])
            self.assertIn('DOCX', rejected['content'][0]['text'])
            self.assertEqual(events, [])
            result = asyncio.run(server['tools'][0].handler({'source_path': 'draft.docx', 'file_name': '天气研究.docx'}))
            receipt = json.loads(result['content'][0]['text'])
            self.assertNotIn('status', receipt)
            self.assertNotIn('pending', str(result))
            self.assertEqual(receipt['name'], '天气研究.docx')
            self.assertEqual(events, [{'type': 'artifact.published', 'artifactId': receipt['artifactId']}])
            self.assertEqual(Document(output / receipt['artifactId'] / 'content').paragraphs[0].text, '天气研究正文')

    def test_publishes_only_from_session_roots(self):
        with tempfile.TemporaryDirectory() as root:
            session = Path(root) / "session"
            work = session / ".work"
            deliverables = session / ".deliverables"
            work.mkdir(parents=True)
            deliverables.mkdir()
            source = work / "draft.py"
            source.write_text("print('internal')", encoding="utf-8")

            result = publish_artifact("draft.py", "最终报告.md", work, deliverables, session)

            self.assertEqual(result["status"], "pending")
            self.assertEqual(result["name"], "最终报告.md")
            snapshot = deliverables / result['artifactId']
            self.assertEqual((snapshot / 'content').read_text(encoding='utf-8'), "print('internal')")
            manifest = json.loads((snapshot / 'manifest.json').read_text(encoding='utf-8'))
            self.assertEqual(manifest['name'], '最终报告.md')
            source.write_text('new version', encoding='utf-8')
            second = publish_artifact('draft.py', '最终报告.md', work, deliverables, session)
            self.assertNotEqual(result['artifactId'], second['artifactId'])
            self.assertEqual((snapshot / 'content').read_text(encoding='utf-8'), "print('internal')")

    def test_rejects_source_outside_session(self):
        with tempfile.TemporaryDirectory() as root:
            session = Path(root) / "session"
            work = session / ".work"
            deliverables = session / ".deliverables"
            work.mkdir(parents=True)
            deliverables.mkdir()
            outside = Path(root) / "outside.txt"
            outside.write_text("secret", encoding="utf-8")

            with self.assertRaisesRegex(ValueError, "当前会话目录"):
                publish_artifact(outside, "secret.txt", work, deliverables, session)

    def test_rejects_directory_traversal_target(self):
        with tempfile.TemporaryDirectory() as root:
            session = Path(root) / "session"
            work = session / ".work"
            deliverables = session / ".deliverables"
            work.mkdir(parents=True)
            deliverables.mkdir()
            source = work / "draft.txt"
            source.write_text("draft", encoding="utf-8")

            with self.assertRaisesRegex(ValueError, "不能包含目录路径"):
                publish_artifact(source, "../secret.txt", work, deliverables, session)


if __name__ == "__main__":
    unittest.main()
