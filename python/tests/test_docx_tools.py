import tempfile
import unittest
from pathlib import Path

from docx import Document

from tools.docx import extract_document, inspect_document, replace_document


class DocxToolsTests(unittest.TestCase):
    def test_inspect_extract_and_replace(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "template.docx"
            document = Document()
            document.add_paragraph("项目：{{project_name}}")
            document.add_table(rows=1, cols=1).cell(0, 0).text = "负责人：{{owner}}"
            document.save(source)

            summary = inspect_document(str(source), root)
            self.assertEqual(summary["placeholders"], ["owner", "project_name"])
            self.assertIn("项目：{{project_name}}", extract_document(str(source), root)["text"])

            result = replace_document(
                str(source), {"project_name": "中文项目", "owner": "李雷"}, root
            )
            self.assertEqual(result["changedParagraphs"], 2)
            self.assertIn("中文项目", extract_document(result["path"], root)["text"])

    def test_paths_cannot_escape_root(self):
        with tempfile.TemporaryDirectory() as directory:
            with self.assertRaises(ValueError):
                inspect_document(str(Path("..") / "outside.docx"), directory)

    def test_uploaded_reference_preserves_reading_order_headings_and_nested_tables(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = root/'reference.docx'
            document = Document()
            document.add_heading('Section', 1)
            table = document.add_table(rows=1, cols=2)
            cell = table.cell(0, 0).merge(table.cell(0, 1))
            cell.text = 'Merged cell'
            cell.add_table(rows=1, cols=1).cell(0, 0).text = 'Nested table'
            document.add_paragraph('After table')
            document.sections[0].header.paragraphs[0].text = 'Header'
            document.save(path)
            summary = inspect_document(str(path), root)
            blocks = summary['structure']
            self.assertEqual([item['type'] for item in blocks], ['paragraph', 'table', 'paragraph'])
            self.assertEqual(blocks[0]['heading_level'], 1)
            self.assertEqual(blocks[1]['rows'][0][0]['column_span'], 2)
            self.assertEqual(blocks[1]['rows'][0][0]['blocks'][1]['type'], 'table')
            text = extract_document(str(path), root)['text']
            self.assertEqual(text.count('Merged cell'), 1)
            self.assertLess(text.index('Nested table'), text.index('After table'))
            self.assertEqual(summary['sections'][0]['header'], ['Header'])
            self.assertNotIn('location_hints', summary)
            self.assertEqual([p.name for p in root.iterdir()], ['reference.docx'])


if __name__ == "__main__":
    unittest.main()
