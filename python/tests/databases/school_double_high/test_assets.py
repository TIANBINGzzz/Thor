"""验证查询资产之间的引用、参数和输出契约，不作为 SQL 安全执行器。"""
from pathlib import Path
import json
import re
import sys
import unittest

ROOT = Path(__file__).resolve().parents[4]
BASE = ROOT / '.claude/databases/schoolDoubleHigh/metrics/hpm'
sys.path.insert(0, str(ROOT/'python'))
from data_access.catalog import Catalog


class AssetContractTests(unittest.TestCase):
    def setUp(self):
        self.catalog = Catalog().freeze(['schoolDoubleHigh'])
        self.specs = {q['id']: self.catalog.spec('schoolDoubleHigh','hpm',q['id'])
                      for q in self.catalog.domain('schoolDoubleHigh','hpm')[1]['queries']}
        self.source = json.loads((BASE.parents[1]/'source.json').read_text(encoding='utf-8'))
        self.tables = self.catalog.schema('schoolDoubleHigh','hpm')['tables']

    def test_sql_parameters_outputs_tables_and_source_registration(self):
        files = set()
        for id, spec in self.specs.items():
            with self.subTest(query=id):
                self.assertEqual(spec['id'],id)
                self.assertTrue(spec['name'] and spec['description'])
                self.assertNotIn('source_keys', spec)
                self.assertEqual(self.source['source_key'], 'schoolDoubleHigh')
                self.assertEqual(len(spec['output']),len({f['name'] for f in spec['output']}))
                for f in spec['output']:
                    self.assertTrue(f['description'])
                    if f['name'].endswith('_id'): self.assertEqual(f['visibility'],'internal_only')
                if spec['sql']:
                    files.add(id)
                    sql=self.catalog.sql('schoolDoubleHigh','hpm',spec)
                    self.assertEqual(set(re.findall(r':([a-z_]+)',sql)),set(spec['parameters']))
                    self.assertTrue(set(re.findall(r'\bt_hpm_[a-z_]+',sql)) <= set(self.tables))
                    self.assertRegex(sql.strip(),r'^(SELECT|WITH)\b')
                    self.assertNotIn(';',sql)  # 一条无尾分号的语句，便于驱动和核验器装配。
                    self.assertNotRegex(sql,r'(?i)SELECT\s+\*|\b(?:INSERT|UPDATE|DELETE|DROP|TRUNCATE|GRANT|OUTFILE|LOAD_FILE|SLEEP)\b')
                else:
                    self.assertIn(spec['status'],('needs_definition','blocked'))
                    self.assertTrue(spec['blockers'])
                for dependency in spec.get('requires_queries',[]) + spec.get('candidate_queries',[]):
                    self.assertIn(dependency,self.specs)
        self.assertEqual(len(files),27)
        self.assertEqual({p.name for p in BASE.iterdir()}, {'projects.yaml','tasks.yaml','performance.yaml','funds.yaml'})

    def test_catalog_and_18_question_coverage(self):
        catalog=self.catalog.domain('schoolDoubleHigh','hpm')[1]
        self.assertEqual({q['id'] for q in catalog['queries']},set(self.specs))
        coverage=json.loads(Path(__file__).with_name('coverage.json').read_text(encoding='utf-8'))
        self.assertEqual(set(coverage['questions']),{f'Q{i}' for i in range(1,19)})
        for mapping in coverage['questions'].values():
            spec=self.specs[mapping['query_id']]
            self.assertEqual(spec['status'],'defined')
            self.assertTrue(set(mapping['bindings']) <= set(spec['parameters']))
            self.assertTrue(mapping['question'])
        documents=self.catalog.documents('schoolDoubleHigh','hpm',list(catalog['documents']))
        self.assertEqual(set(re.findall(r'^## (t_hpm_\w+)$',documents['schema'],re.M)),set(self.tables))
        for entity in catalog['entities'].values():
            output={f['name'] for f in self.specs[entity['query_id']]['output']}
            self.assertTrue({entity['id_field'],entity['name_field']} <= output)

    def test_schema_has_verified_types_and_no_duplicate_workflow_binding(self):
        self.assertEqual(len(self.tables),13)
        self.assertEqual(sum(map(len,self.tables.values())),158)
        self.assertNotIn('UNKNOWN',{kind for table in self.tables.values() for kind in table.values()})
        for workflow in ('double-high-qa','writing-docx'):
            config=json.loads((ROOT/f'.claude/workflows/{workflow}/workflow.json').read_text(encoding='utf-8'))
            self.assertNotIn('data_sources',config)
        self.assertEqual(self.source['capabilities'],['national-excellence-data-qa','document-writing'])

    def test_template_guide_references_metrics_and_source_keeps_all_tables(self):
        from docx import Document
        from workflows.writing_docx.template_assets import load_template
        root=ROOT/'.claude/workflows/writing-docx/templates/szpt-midterm'
        guide=(root/'writing-guide.md').read_text(encoding='utf-8')
        referenced=set(re.findall(r'`([a-z][a-z0-9_]+)`',guide))
        self.assertTrue(referenced)
        self.assertTrue(referenced <= set(self.specs) | {'szpt_midterm'})
        template=load_template('szpt-midterm','document-writing')
        self.assertEqual(set(template['assets']),{'docx','writing_guide'})
        self.assertEqual(len(Document(template['_docx']).tables),40)


if __name__ == '__main__':
    unittest.main()
