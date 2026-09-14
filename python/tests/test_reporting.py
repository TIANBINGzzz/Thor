import asyncio
from copy import deepcopy
from hashlib import sha256
import json
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile
import unittest

from docx import Document

from data_access.context import DataError
from data_access.runtime import RunServices
from reporting.bindings import load_template, document_xml, text_hash, NS
from reporting.planner import Planner
from reporting.rendering import render
from tests import test_data_access as fixture
save = fixture.save


class ReportingTests(unittest.TestCase):
    setUp = fixture.DataAccessTests.setUp
    make_executor = fixture.DataAccessTests.make_executor
    school = fixture.DataAccessTests.school

    def template(self, missing=False):
        directory=self.root/'templates/demo'
        directory.mkdir(parents=True)
        document=Document()
        document.add_heading('Report',0)
        table=document.add_table(rows=1, cols=2)
        table.cell(0,0).text='Total'
        table.cell(0,1).text='pending'
        document.add_paragraph('Pending scoring' if missing else 'Static footer')
        path=directory/'template.docx'
        document.save(path)
        xml=document_xml(path)
        paragraphs=xml.xpath('.//w:p',namespaces=NS)
        slots=[]
        for index,node in enumerate(paragraphs):
            dynamic=index==2 or (missing and index==3)
            slot={'slot_key':f's{index}','section_key':'summary','name':f'Field {index}',
                'kind':'scalar' if dynamic else 'static','required':dynamic,
                'locator':{'part':'word/document.xml','path':xml.getroottree().getpath(node),'expected_text_hash':text_hash(node)}}
            if index==2:
                slot.update(dataset_key='total',field='total',format={'unit':'USD','decimal_places':2})
            slots.append(slot)
        save(directory/'slots.json',{'slots':slots})
        save(directory/'bindings.json',{'template_version':1,'slot_files':['slots.json'],'datasets':[
            {'dataset_key':'total','source_role':'sales','domain':'sales','scope_role':'school',
             'query_id':'total','parameter_bindings':{},'section_keys':['summary']} ]})
        save(directory/'template.json',{'template_key':'demo','enabled':True,'capabilities':['qa'],
            'version':1,'docx_file':path.name,'docx_sha256':sha256(path.read_bytes()).hexdigest(),
            'bindings_file':'bindings.json','source_roles':{'sales':'first'},'scope_roles':{'sales':['school']},
            'missing_policy':'annotated_draft','parameters':{'type':'object'}})
        return load_template('demo','qa',directory.parent)

    def test_plan_deduplicates_snapshot_and_formats_fixed_values(self):
        async def scenario():
            template=self.template()
            duplicate=deepcopy(template['_bindings']['datasets'][0])
            duplicate['dataset_key']='repeated_total'
            template['_bindings']['datasets'].append(duplicate)
            services=RunServices(['first','second'],env={},catalog=self.catalog)
            services.executor=self.executor
            planner=Planner(services,template)
            services.report=planner
            params={'years':['2025'],'period_mode':'annual','timezone':'Asia/Shanghai'}
            scopes={'sales':{'school':self.school()}}
            result=await planner.prepare(params,scopes)
            repeated=await planner.prepare(params,scopes)
            self.assertEqual(result['plan_ref'],repeated['plan_ref'])
            await asyncio.gather(*list(planner.tasks))
            plan=planner.plans[result['plan_ref']]
            self.assertEqual(plan['status'],'ready')
            self.assertEqual(len(plan['nodes']),1)
            output=await planner.render(result['plan_ref'],[])
            self.assertEqual(output['status'],'complete')
            self.assertIn('10.00',Document(output['path']).tables[0].cell(0,1).text)
            with ZipFile(template['_docx']) as old, ZipFile(output['path']) as new:
                self.assertEqual(old.namelist(),new.namelist())
                for name in old.namelist():
                    if name!='word/document.xml': self.assertEqual(old.read(name),new.read(name))
            await services.close()
        asyncio.run(scenario())

    def test_missing_definitions_produce_annotated_draft_and_reject_invented_evidence(self):
        async def scenario():
            template=self.template(missing=True)
            services=RunServices(['first'],env={},catalog=self.catalog)
            services.executor=self.executor
            planner=Planner(services,template)
            services.report=planner
            result=await planner.prepare({'years':['2025'],'period_mode':'annual','timezone':'Asia/Shanghai'},
                {'sales':{'school':self.school()}})
            await asyncio.gather(*list(planner.tasks))
            data=planner.get(result['plan_ref'],'summary')
            self.assertEqual(data['status'],'blocked')
            self.assertEqual(data['coverage']['required_missing'],1)
            output=await planner.render(result['plan_ref'],[])
            self.assertEqual(output['status'],'draft')
            self.assertEqual(output['missing_count'],1)
            self.assertIn('待核验',Document(output['path']).paragraphs[-1].text)
            with self.assertRaises(DataError):
                await planner.render(result['plan_ref'],[{'section_key':'summary','text':'999','evidence_refs':['invented']}])
            with self.assertRaises(DataError): planner.get('other-plan')
            await services.close()
        asyncio.run(scenario())

    def test_template_drift_and_duplicate_locations_fail(self):
        template=self.template()
        path=template['_directory']/'slots.json'
        slots=json.loads(path.read_text())
        slots['slots'].append(slots['slots'][0])
        save(path,slots)
        with self.assertRaisesRegex(DataError,'DUPLICATE_SLOT'): load_template('demo','qa',path.parent.parent)
        document=Document(template['_docx'])
        document.add_paragraph('changed')
        document.save(template['_docx'])
        with self.assertRaisesRegex(DataError,'TEMPLATE_MISMATCH'): load_template('demo','qa',path.parent.parent)

    def test_full_template_draft_preserves_all_40_tables_and_package_parts(self):
        template=load_template('szpt-midterm','document-writing')
        executor=self.make_executor()
        from dataclasses import replace
        executor.context=replace(executor.context,capability_ref='document-writing')
        slots=[{**s,'value_status':'static' if s['kind']=='static' else 'definition_missing'} for s in template['_slots']]
        output=render(template,{'plan_ref':'test-draft','slots':slots},[],executor)
        self.assertEqual(output['status'],'draft')
        before=document_xml(template['_docx'])
        after=document_xml(output['path'])
        for tag in ('tbl','tr','tc','p','sectPr','drawing'):
            self.assertEqual(len(before.xpath(f'.//w:{tag}',namespaces=NS)),len(after.xpath(f'.//w:{tag}',namespaces=NS)))
        with ZipFile(template['_docx']) as old, ZipFile(output['path']) as new:
            for name in old.namelist():
                if name!='word/document.xml': self.assertEqual(old.read(name),new.read(name))


if __name__ == '__main__': unittest.main()
