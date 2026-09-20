"""平台文件下载的真实响应形态、凭据边界和失败清理。"""

import asyncio
import hashlib
import io
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import httpx
from docx import Document

from runtime.file_broker import FileBroker, FileBrokerError, FileBrokerConfigurationError
from runtime.protocol import AttachmentRef


class FileServiceDownloadTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        (self.root / 'input').mkdir()
        self.config = self.root / 'databases.json'
        self.settings = {'version': 1, 'sources': {}, 'fileService': {
            'baseUrl': 'https://files.example.test',
            'domainName': 'routing.example.test',
            'remoteUrl': 'https://public.example.test',
            'downloadPath': '/fwk_manage_service/sys_attachment/{fileId}/ai/download/',
        }}
        self.save()
        env = patch.dict('os.environ', {'CCSDK_DATABASES_FILE': str(self.config)}, clear=True)
        env.start()
        self.addCleanup(env.stop)

    def save(self):
        self.config.write_text(json.dumps(self.settings), encoding='utf-8')

    async def fetch(self, handler, refs=None, **kwargs):
        broker = FileBroker(transport=httpx.MockTransport(handler), **kwargs)
        return await broker.fetch_all(refs or [AttachmentRef('file_1', 'reference')],
            run_id='run_1', tenant_id='tenant', user_id='user',
            workspace=self.root / 'input', bearer_token='private-run-jwt')

    async def test_downloads_docx_using_configured_route_without_forwarding_credentials(self):
        document = Document()
        document.add_heading('自定义模板', level=1)
        stream = io.BytesIO()
        document.save(stream)
        content = stream.getvalue()
        requests, events = [], []

        def handler(request):
            requests.append(request)
            return httpx.Response(200, content=content, headers={
                'content-type': 'application/octet-stream',
                'content-disposition': 'attachment;filename=%E6%A8%A1%E6%9D%BF+1.docx'})

        files = await self.fetch(handler, event_sink=lambda kind, data: events.append((kind, data)))
        self.assertEqual(len(requests), 1)
        request = requests[0]
        self.assertEqual(request.method, 'GET')
        self.assertEqual(str(request.url),
            'https://files.example.test/fwk_manage_service/sys_attachment/file_1/ai/download/')
        self.assertEqual(request.headers['domain-name'], 'routing.example.test')
        self.assertEqual(request.headers['remote-url'], 'https://public.example.test')
        self.assertNotIn('authorization', request.headers)
        self.assertEqual(request.content, b'')
        self.assertEqual(files[0].safe_name, '模板 1.docx')
        self.assertEqual(Document(files[0].path).paragraphs[0].text, '自定义模板')
        self.assertEqual(files[0].sha256, hashlib.sha256(content).hexdigest())
        self.assertEqual(files[0].to_metadata()['source'], 'file-service')
        self.assertNotIn('private-run-jwt', str(events))
        self.assertNotIn('example.test', str(events))

    async def test_chunked_pdf_without_remote_digest_and_rfc_filename(self):
        class Stream(httpx.AsyncByteStream):
            async def __aiter__(self):
                yield b'%PDF-'
                yield b'content'
        files = await self.fetch(lambda _: httpx.Response(200, stream=Stream(), headers={
            'content-type': 'image/pdf',
            'content-disposition': "attachment; filename*=UTF-8''%E6%8A%A5%E5%91%8A%2B1.pdf"}))
        self.assertEqual(files[0].safe_name, '报告+1.pdf')
        self.assertEqual(files[0].path.read_bytes(), b'%PDF-content')
        self.assertEqual(files[0].size, 12)

    async def test_platform_download_does_not_load_broker_credentials_or_reuse_cookies(self):
        requests = []
        def handler(request):
            requests.append(request)
            return httpx.Response(200, content=b'%PDF-content', headers={
                'content-disposition': 'attachment;filename=source.pdf',
                'set-cookie': 'broker-session=private; Path=/; Secure'})
        with patch.dict('os.environ', {
            'CCSDK_FILE_BROKER_AUTH_MODE': 'service',
            'CCSDK_FILE_BROKER_SERVICE_TOKEN': 'private-service-token',
            'CCSDK_FILE_BROKER_CA': 'missing-broker-ca.pem',
            'CCSDK_FILE_BROKER_CLIENT_CERT': 'missing-broker-cert.pem',
            'CCSDK_FILE_BROKER_CLIENT_KEY': 'missing-broker-key.pem',
        }), patch('runtime.file_broker.httpx.AsyncClient', wraps=httpx.AsyncClient) as client:
            files = await self.fetch(handler, [AttachmentRef('file_1'), AttachmentRef('file_2')])
        self.assertIsNone(client.call_args.kwargs['cert'])
        self.assertIs(client.call_args.kwargs['verify'], True)
        self.assertEqual(len(files), 2)
        self.assertNotEqual(files[0].path, files[1].path)
        for request in requests:
            self.assertNotIn('cookie', request.headers)
            self.assertNotIn('authorization', request.headers)

    async def test_download_rejects_redirect_errors_and_unsafe_filename(self):
        replies = [
            httpx.Response(302, headers={'location': 'https://other.test/file'}),
            httpx.Response(403), httpx.Response(404), httpx.Response(500),
            httpx.Response(200, json={'state': 500}),
            httpx.Response(200, text='<html>login</html>', headers={'content-type': 'text/html'}),
            httpx.Response(200, content=b'bad', headers={
                'content-disposition': 'attachment;filename=..%2Fsecret.docx'}),
        ]
        for reply in replies:
            with self.subTest(status=reply.status_code, headers=dict(reply.headers)):
                with self.assertRaises(FileBrokerError):
                    await self.fetch(lambda _: reply)
                self.assertEqual(list((self.root / 'input').iterdir()), [])

    async def test_size_limit_and_broken_stream_remove_partial_and_prior_files(self):
        for outcome in ('limit', 'broken', 'cancelled', 'truncated'):
            class Stream(httpx.AsyncByteStream):
                async def __aiter__(self):
                    yield b'x' * 262144
                    if outcome == 'broken':
                        raise httpx.ReadError('private upstream address')
                    if outcome == 'cancelled':
                        raise asyncio.CancelledError()
                    if outcome == 'limit':
                        yield b'x' * 262144

            def handler(request):
                headers = {'content-disposition': 'attachment;filename=source.pdf',
                           'content-type': 'application/pdf'}
                if '/file_1/' in request.url.path:
                    return httpx.Response(200, content=b'%PDF-first', headers=headers)
                if outcome == 'truncated':
                    headers['content-length'] = '300000'
                return httpx.Response(200, headers=headers, stream=Stream())

            with self.subTest(outcome=outcome):
                expected = asyncio.CancelledError if outcome == 'cancelled' else FileBrokerError
                with self.assertRaises(expected):
                    await self.fetch(handler, [AttachmentRef('file_1'), AttachmentRef('file_2')], max_bytes=300000)
                self.assertEqual(list((self.root / 'input').iterdir()), [])

    async def test_invalid_ids_never_reach_network(self):
        for file_id in ('../other', 'a/b', 'a?query=1', 'a#fragment', ''):
            with self.subTest(file_id=file_id), self.assertRaises(FileBrokerError):
                await self.fetch(lambda _: self.fail('unsafe file ID sent'), [AttachmentRef(file_id)])

    async def test_writing_run_prepares_external_template_before_sdk_and_cleans_input(self):
        import server
        from runtime.protocol import AgentRunRequest
        from runtime.run_store import RunStore
        from runtime.session_actor import SessionManager
        from tests.test_session_actor import FakeClientWorker, complete_script

        document = Document()
        document.add_heading('会议纪要模板', level=1)
        stream = io.BytesIO()
        document.save(stream)
        for mode in ('query', 'client'):
            with self.subTest(mode=mode):
                store = RunStore(':memory:')
                self.addCleanup(store.close)
                request = AgentRunRequest.from_dict({'protocol': 'agent-run/v1', 'runId': 'run_1',
                    'messageId': 'msg_1', 'capabilityRef': 'document-writing',
                    'input': {'text': '根据模板撰写会议纪要，仅使用提供的材料。',
                              'attachmentRefs': [{'fileId': 'file_1', 'purpose': 'reference'}]}})
                store.create_run('run_1', request=request, tenant_id='tenant', user_id='user')
                sdk_calls, fetched = [], []

                def assert_ready(payload):
                    sdk_calls.append(payload)
                    self.assertIn('template.docx', payload['prompt'])
                    paths = list(self.root.rglob('template.docx'))
                    self.assertEqual(len(paths), 1)
                    self.assertEqual(Document(paths[0]).paragraphs[0].text, '会议纪要模板')

                async def query(payload):
                    assert_ready(payload)
                    yield {'type': 'result', 'ok': True}

                class Worker(FakeClientWorker):
                    async def send(worker, message):
                        if message['type'] == 'client_query':
                            assert_ready(message)
                        await super().send(message)

                manager = SessionManager(worker_factory=lambda **_: Worker([complete_script()], worker_id=0))
                self.addAsyncCleanup(manager.close_all)

                def response(request):
                    fetched.append(request)
                    return httpx.Response(200, content=stream.getvalue(), headers={
                        'content-disposition': 'attachment;filename=template.docx',
                        'content-type': 'application/octet-stream'})

                def broker(**kwargs):
                    return FileBroker(transport=httpx.MockTransport(response), **kwargs)

                with patch.multiple(server, RUN_STORE=store, PROJECT_ROOT=self.root,
                        CLIENT_SESSION_ROOT=self.root / 'sessions', SESSION_MANAGER=manager,
                        MODELS=['test'], internal_tasks={}, internal_subscribers={}), \
                     patch.object(server, '_runtime_mode_for_request', return_value=mode), \
                     patch.object(server, 'FileBroker', side_effect=broker), \
                     patch.object(server, 'stream_agent', query):
                    await server._execute_internal_run(request, {'tenant': 'tenant', 'sub': 'user'},
                                                       runtime_bearer='private-run-jwt')
                    self.assertEqual(store.get_run('run_1')['status'], 'succeeded', store.get_run('run_1'))
                    self.assertEqual(len(sdk_calls), 1)
                    self.assertEqual(fetched[0].method, 'GET')
                    self.assertFalse(list(self.root.rglob('template.docx')))
                    events = store.events_after('run_1', 0)
                    phases = [e['payload']['name'] for e in events if e['type'] == 'phase']
                    self.assertLess(phases.index('files_ready'), phases.index('model_starting'))
                    for secret in ('private-run-jwt', 'files.example.test', str(self.root)):
                        self.assertNotIn(secret, json.dumps(events))
                    await manager.close_all()

    def test_invalid_download_configuration_is_rejected(self):
        for field, value in (
            ('downloadPath', 'https://other.test/{fileId}'),
            ('downloadPath', '/file/{fileId}/{unknown}'),
            ('downloadPath', '/file/{fileId}?extra=1'),
            ('downloadPath', '/file/static'),
            ('downloadPath', '/file/../{fileId}'),
            ('domainName', 'host\r\ninjected: true'),
            ('baseUrl', 'https://files.test:invalid'),
        ):
            with self.subTest(field=field, value=value):
                original = self.settings['fileService'][field]
                self.settings['fileService'][field] = value
                self.save()
                with self.assertRaises(FileBrokerConfigurationError):
                    FileBroker()
                self.settings['fileService'][field] = original
