import tempfile
import subprocess
import zipfile
import sys
import threading
import time
import unittest
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from tools.document_conversion import _prepare_fonts, _run_container, read_pdf, render_document


def make_pdf(path):
    streams = [b"BT /F1 18 Tf 30 100 Td (First page) Tj ET", b"BT /F1 18 Tf 30 100 Td (Second page) Tj ET"]
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R 4 0 R] /Count 2 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 240 160] /Resources << /Font << /F1 5 0 R >> >> /Contents 6 0 R >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 240 160] /Resources << /Font << /F1 5 0 R >> >> /Contents 7 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        *[b"<< /Length " + str(len(s)).encode() + b" >>\nstream\n" + s + b"\nendstream" for s in streams],
    ]
    data, offsets = b"%PDF-1.4\n", [0]
    for index, obj in enumerate(objects, 1):
        offsets.append(len(data))
        data += f"{index} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref = len(data)
    data += b"xref\n0 8\n0000000000 65535 f \n"
    data += b"".join(f"{n:010} 00000 n \n".encode() for n in offsets[1:])
    data += f"trailer\n<< /Size 8 /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    path.write_bytes(data)


class DocumentConversionTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.source = self.root / "source.docx"
        document = Document()
        document.add_heading("Conversion sample", 1)
        document.add_paragraph("Document body")
        document.save(self.source)

    def test_subset_fonts_are_detached_only_in_render_copy(self):
        from lxml import etree
        with zipfile.ZipFile(self.source) as archive:
            parts = {item.filename: archive.read(item.filename) for item in archive.infolist()}
        fonts = etree.fromstring(parts["word/fontTable.xml"])
        subset = etree.SubElement(fonts[0], qn("w:embedRegular"))
        subset.set(qn("w:subsetted"), "1")
        complete = etree.SubElement(fonts[0], qn("w:embedBold"))
        complete.set(qn("w:subsetted"), "0")
        parts["word/fontTable.xml"] = etree.tostring(fonts)
        with zipfile.ZipFile(self.source, "w") as archive:
            for name, data in parts.items():
                archive.writestr(name, data)
        original = self.source.read_bytes()
        prepared, names = _prepare_fonts(self.source, self.root)
        self.assertEqual(names, [fonts[0].get(qn("w:name"))])
        self.assertNotEqual(prepared, self.source)
        self.assertEqual(self.source.read_bytes(), original)
        with zipfile.ZipFile(prepared) as archive:
            after = etree.fromstring(archive.read("word/fontTable.xml"))
            self.assertIsNone(after[0].find(qn("w:embedRegular")))
            self.assertIsNotNone(after[0].find(qn("w:embedBold")))
            for name, content in parts.items():
                if name != "word/fontTable.xml":
                    self.assertEqual(archive.read(name), content)

    def test_container_timeout_removes_only_its_own_instance(self):
        def run(command, **kwargs):
            if command[1] == "run":
                raise subprocess.TimeoutExpired(command, 180)
            return subprocess.CompletedProcess(command, 0)
        with patch("tools.document_conversion.shutil.which", return_value="docker"), \
                patch("tools.document_conversion.subprocess.run", side_effect=run) as process:
            with self.assertRaisesRegex(RuntimeError, "超时"):
                _run_container(self.source, self.root / "updated.docx", self.root / "rendered.pdf", self.root)
        command = process.call_args_list[0].args[0]
        self.assertIn("--network=none", command)
        self.assertIn(f"type=bind,source={self.source},target=/input.docx,readonly", command)
        self.assertEqual(process.call_args_list[1].args[0],
                         ["docker", "rm", "--force", command[command.index("--name") + 1]])

    def test_pdf_read_in_fresh_worker_thread_does_not_load_numpy(self):
        make_pdf(self.root / "sample.pdf")
        script = """import asyncio, sys
from tools.document_conversion import read_pdf
async def main():
    result = await asyncio.to_thread(read_pdf, sys.argv[1], sys.argv[2], render=True)
    assert result['pageCount'] == 2
    assert 'numpy' not in sys.modules
asyncio.run(main())
"""
        result = subprocess.run([sys.executable, "-c", script, str(self.root / "sample.pdf"), str(self.root)],
                                cwd=Path(__file__).resolve().parents[1],
                                capture_output=True, text=True, timeout=20)
        self.assertEqual(result.returncode, 0, result.stderr)

    def test_read_pdf_paginates_text_and_renders_only_selected_page(self):
        make_pdf(self.root / "sample.pdf")
        result = read_pdf("sample.pdf", self.root, start=2, limit=1, render=True, output_dir="pages")
        self.assertEqual(result["pageCount"], 2)
        self.assertIsNone(result["nextStart"])
        self.assertEqual(len(result["pages"]), 1)
        page = result["pages"][0]
        self.assertEqual(page["page"], 2)
        self.assertIn("Second page", page["text"])
        self.assertTrue(Path(page["imagePath"]).read_bytes().startswith(b"\x89PNG"))
        self.assertEqual(len(list((self.root / "pages").glob("*.png"))), 1)

    def test_pdf_rejects_invalid_ranges_and_output_escape(self):
        make_pdf(self.root / "sample.pdf")
        for args in [{"start": 0}, {"start": 3}, {"limit": 0}, {"limit": 21}, {"render": True, "output_dir": "../pages"}]:
            with self.subTest(args=args), self.assertRaises(ValueError):
                read_pdf("sample.pdf", self.root, **args)

    def test_parallel_pdf_calls_serialize_pdfium_lifetime(self):
        import pypdfium2

        make_pdf(self.root / "sample.pdf")
        original = pypdfium2.PdfDocument
        active, maximum = 0, 0
        lock = threading.Lock()

        def open_document(*args, **kwargs):
            nonlocal active, maximum
            with lock:
                active += 1
                maximum = max(maximum, active)
            time.sleep(0.03)
            document = original(*args, **kwargs)
            close = document.close

            def finish():
                nonlocal active
                try:
                    close()
                finally:
                    with lock:
                        active -= 1
            document.close = finish
            return document

        with patch("pypdfium2.PdfDocument", side_effect=open_document), ThreadPoolExecutor(max_workers=2) as pool:
            results = list(pool.map(lambda _: read_pdf("sample.pdf", self.root), range(2)))
        self.assertEqual([result["pageCount"] for result in results], [2, 2])
        self.assertEqual(maximum, 1)

    def test_additional_directory_is_allowed_but_symlink_escape_is_not(self):
        work = self.root / "work"
        work.mkdir()
        make_pdf(self.root / "sample.pdf")
        result = read_pdf(str(self.root / "sample.pdf"), work, [str(self.root)])
        self.assertEqual(result["nextStart"], None)
        with self.assertRaises(ValueError):
            read_pdf(str(self.root / "sample.pdf"), work)
        link = work / "linked.pdf"
        try:
            link.symlink_to(self.root / "sample.pdf")
        except OSError:
            return
        with self.assertRaises(ValueError):
            read_pdf("linked.pdf", work)

    def test_render_requires_actual_toc_refresh_and_preserves_source(self):
        document = Document(self.source)
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), 'TOC \\o "1-3"')
        document.add_paragraph()._p.append(field)
        document.save(self.source)
        original = self.source.read_bytes()

        def office(source, docx_path, pdf_path, work_dir):
            docx_path.write_bytes(source.read_bytes())
            make_pdf(pdf_path)
            return {"engine": "test-office", "indexCount": 1, "tocCount": 1, "indexesUpdated": 1, "fieldsUpdated": True}

        with patch("tools.document_conversion._run_office", side_effect=office):
            result = render_document("source.docx", "rendered", self.root)
        self.assertEqual(result["tocStatus"], "updated")
        self.assertEqual(result["engine"], "test-office")
        self.assertTrue(Path(result["docxPath"]).is_file())
        self.assertTrue(Path(result["pdfPath"]).is_file())
        self.assertEqual(self.source.read_bytes(), original)
        with patch("tools.document_conversion._run_office", side_effect=office), self.assertRaises(FileExistsError):
            render_document("source.docx", "rendered", self.root)

    def test_render_rejects_missing_refresh_and_does_not_publish_partial_files(self):
        def office(source, docx_path, pdf_path, work_dir):
            docx_path.write_bytes(source.read_bytes())
            make_pdf(pdf_path)
            return {"engine": "test-office", "indexCount": 1, "tocCount": 1, "indexesUpdated": 0, "fieldsUpdated": True}

        with patch("tools.document_conversion._run_office", side_effect=office), self.assertRaises(RuntimeError):
            render_document("source.docx", "rendered", self.root)
        self.assertEqual(list((self.root / "rendered").glob("*.docx")), [])
        self.assertEqual(list((self.root / "rendered").glob("*.pdf")), [])

    def test_render_rejects_source_toc_disappearing_in_office(self):
        document = Document(self.source)
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), 'TOC \\o "1-3"')
        document.add_paragraph()._p.append(field)
        document.save(self.source)

        def office(source, docx_path, pdf_path, work_dir):
            docx_path.write_bytes(source.read_bytes())
            make_pdf(pdf_path)
            return {"engine": "test-office", "indexCount": 0, "tocCount": 0, "indexesUpdated": 0, "fieldsUpdated": True}

        with patch("tools.document_conversion._run_office", side_effect=office), self.assertRaises(RuntimeError):
            render_document("source.docx", "rendered", self.root)
        self.assertEqual(list((self.root / "rendered").glob("*.pdf")), [])


if __name__ == "__main__":
    unittest.main()
