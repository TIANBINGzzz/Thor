"""实际上游二进制回归：批量回写保留合并结构，失败批次不留下半成品。"""

import json
import os
from pathlib import Path
import shutil
import subprocess
import tempfile
import unittest

from docx import Document


OFFICECLI = os.environ.get('CCSDK_OFFICECLI_PATH') or shutil.which('officecli')


@unittest.skipUnless(OFFICECLI, '需安装 OfficeCLI 或设置 CCSDK_OFFICECLI_PATH')
class OfficeCLIIntegrationTests(unittest.TestCase):
    def test_batch_preserves_template_and_rolls_back_invalid_operation(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source, draft = root / 'source.docx', root / 'draft.docx'
            document = Document()
            document.add_heading('原题名', 1)
            paragraph = document.add_paragraph()
            paragraph.add_run('保留粗体').bold = True
            paragraph.add_run('保留斜体').italic = True
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).merge(table.cell(0, 1)).text = '合并表头'
            table.cell(1, 0).text = '指标'
            table.cell(1, 1).text = '待填写'
            document.save(source)
            original = source.read_bytes()
            shutil.copyfile(source, draft)

            def call(*arguments):
                return subprocess.run([OFFICECLI, *map(str, arguments)], capture_output=True,
                                      text=True, encoding='utf-8', timeout=45,
                                      env={**os.environ, 'OFFICECLI_SKIP_UPDATE': '1'})

            try:
                edits = root / 'edits.json'
                edits.write_text(json.dumps([
                    {'command': 'set', 'path': '/body/p[1]', 'props': {'text': '新题名'}},
                    {'command': 'set', 'path': '/body/tbl[1]/tr[2]/tc[2]', 'props': {'text': '12'}},
                ], ensure_ascii=False), encoding='utf-8')
                result = call('batch', draft, '--input', edits, '--json')
                self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
                self.assertTrue(json.loads(result.stdout)['success'])
                self.assertEqual(call('close', draft).returncode, 0)
                updated = Document(draft)
                self.assertEqual(updated.paragraphs[0].text, '新题名')
                self.assertEqual(updated.paragraphs[0].style.name, 'Heading 1')
                self.assertTrue(updated.paragraphs[1].runs[0].bold)
                self.assertTrue(updated.paragraphs[1].runs[1].italic)
                self.assertEqual(updated.tables[0].cell(1, 1).text, '12')
                self.assertEqual(updated.tables[0].cell(0, 0)._tc.grid_span, 2)

                edits.write_text(json.dumps([
                    {'command': 'set', 'path': '/body/p[1]', 'props': {'text': '不应保存'}},
                    {'command': 'set', 'path': '/body/tbl[99]', 'props': {'text': '无效'}},
                ], ensure_ascii=False), encoding='utf-8')
                call('batch', draft, '--input', edits, '--json')
                self.assertEqual(call('close', draft).returncode, 0)
                self.assertEqual(Document(draft).paragraphs[0].text, '新题名')
                self.assertEqual(source.read_bytes(), original)
            finally:
                call('close', draft)
