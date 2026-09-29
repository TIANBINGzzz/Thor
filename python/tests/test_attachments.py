import base64
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document

from runtime.capabilities import CAPABILITIES
from runtime.config import build_options
from tools.attachments import read_attachment


class AttachmentTests(unittest.TestCase):
    def test_every_capability_mounts_the_same_scoped_reader(self):
        with tempfile.TemporaryDirectory() as folder:
            with patch('runtime.nacos_config.fetch_config', return_value={
                    'campusMcp': {
                        'url': 'https://campus.example.test/string_campus_brain_service/mcp',
                        'domainName': 'campus.example.test'}}):
                for capability in CAPABILITIES.values():
                    with self.subTest(capability=capability.ref), patch.dict('os.environ', {
                            'CCSDK_IMAGE_BASE_URL': 'https://model.test', 'CCSDK_IMAGE_API_KEY': 'test-key'}):
                        self.assertTrue(capability.supports_attachments)
                        options = build_options({'capability_ref': capability.ref,
                            'workflow_name': capability.workflow_ref, 'input_directory': folder,
                            'credentials': {'platformBearer': 'test-token'}})
                        self.assertIn('attachments', options.mcp_servers)
                        self.assertIn('mcp__attachments__read', options.allowed_tools)

    def test_text_pagination_docx_tables_and_path_boundary(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder) / 'input'
            root.mkdir()
            (root / 'data.csv').write_text('学校,经费\n' + '甲' * 13000, encoding='utf-8')
            first = json.loads(read_attachment(root, 'data.csv')['content'][0]['text'])
            second = json.loads(read_attachment(root, 'data.csv', first['next_offset'])['content'][0]['text'])
            self.assertEqual(first['text'] + second['text'], (root / 'data.csv').read_text(encoding='utf-8'))
            self.assertIsNone(second['next_offset'])
            document = Document()
            document.add_paragraph('附件说明')
            document.add_table(rows=1, cols=1).cell(0, 0).text = '表格数据'
            document.save(root / 'data.docx')
            text = json.loads(read_attachment(root, 'data.docx')['content'][0]['text'])['text']
            self.assertIn('附件说明', text)
            self.assertIn('表格数据', text)
            outside = Path(folder) / 'secret.txt'
            outside.write_text('private')
            for path in (str(outside), '../secret.txt'):
                with self.assertRaises(ValueError):
                    read_attachment(root, path)

    def test_pdf_and_image_use_content_not_unrestricted_read(self):
        import pypdfium2 as pdfium
        from PIL import Image
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            document = pdfium.PdfDocument.new()
            page = document.new_page(100, 100)
            document.save(root / 'page.pdf')
            page.close()
            document.close()
            result = json.loads(read_attachment(root, 'page.pdf')['content'][0]['text'])
            self.assertEqual(result['pageCount'], 1)
            self.assertEqual(result['pages'][0]['page'], 1)
            Image.new('RGB', (2, 2)).save(root / 'image.png')
            image = read_attachment(root, 'image.png')['content'][0]
            self.assertEqual(image['type'], 'image')
            self.assertEqual(base64.b64decode(image['data']), (root / 'image.png').read_bytes())
