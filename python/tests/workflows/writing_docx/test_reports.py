import asyncio
from copy import deepcopy
from hashlib import sha256
import json
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile
import unittest

from docx import Document

from data_access.context import DataError
from runtime.data_services import RunServices
from workflows.writing_docx.document_map import build_document_map, document_xml, text_hash, NS
from workflows.writing_docx.template_assets import load_template
from workflows.writing_docx.report_planner import ReportPlanner
from workflows.writing_docx.report_renderer import render
from workflows.writing_docx.report_validator import check_fidelity, validated_plan, record_page_review
from tests import test_data_access as fixture
save = fixture.save


class ReportingTests(unittest.TestCase):
    setUp = fixture.DataAccessTests.setUp
    make_executor = fixture.DataAccessTests.make_executor
    save_connections = fixture.DataAccessTests.save_connections
    school = fixture.DataAccessTests.school

    def template(self, missing=False):
        directory=self.root/'templates/demo'
        directory.mkdir(parents=True)
        document=Document()
        document.add_heading('Report',0)
        table=document.add_table(rows=1, cols=2)
        table.cell(0,0).text='Total'
        table.cell(0,1).text='pending'
        document.add_paragraph('Pending scoring' if missing else 'Static footer')
        path=directory/'template.docx'
        document.save(path)
        xml=document_xml(path)
        paragraphs=xml.xpath('.//w:p',namespaces=NS)
        slots=[]
        for index,node in enumerate(paragraphs):
            dynamic=index==2 or (missing and index==3)
            slot={'slot_key':f's{index}','section_key':'summary','name':f'Field {index}',
                'kind':'scalar' if dynamic else 'static','required':dynamic,
                'locator':{'part':'word/document.xml','path':xml.getroottree().getpath(node),'expected_text_hash':text_hash(node)}}
            if index==2:
                slot.update(text_template='${amount}',values={'amount':{'dataset_key':'total','field':'total'}})
            slots.append(slot)
        save(directory/'document-map.json',build_document_map('demo',1,path,slots))
        save(directory/'report-data-plan.json',{'plan_version':1,'template_key':'demo','template_version':1,'datasets':[
            {'dataset_key':'total','source_role':'sales','domain':'sales','scope_role':'school',
             'query_id':'total','parameter_bindings':{}} ]})
        save(directory/'template.json',{'template_key':'demo','enabled':True,'capabilities':['qa'],
            'version':1,'assets':{'docx':{'file':path.name,'sha256':sha256(path.read_bytes()).hexdigest()},
            'writing_guide':'writing-guide.md','document_map':'document-map.json','data_plan':'report-data-plan.json'},
            'data':{'source_roles':{'sales':'first'},'scope_roles':{'sales':['school']}},
            'report':{'parameters':{'type':'object'},'file_name':'report.docx'},
            'output_policy':{'missing_policy':'reject','preserve_structure':True}})
        (directory/'writing-guide.md').write_text('DEMO GUIDE',encoding='utf-8')
        return load_template('demo','qa',directory.parent)

    def test_plan_deduplicates_snapshot_and_fills_fixed_values(self):
        async def scenario():
            template=self.template()
            duplicate=deepcopy(template['_bindings']['datasets'][0])
            duplicate['dataset_key']='repeated_total'
            template['_bindings']['datasets'].append(duplicate)
            services=RunServices(['first','second'],env=self.env,catalog=self.catalog)
            services.executor=self.executor
            planner=ReportPlanner(services,template)
            services.report=planner
            params={'years':['2025'],'period_mode':'annual','timezone':'Asia/Shanghai'}
            scopes={'sales':{'school':self.school()}}
            result=await planner.prepare(params,scopes)
            repeated=await planner.prepare(params,scopes)
            self.assertEqual(result['plan_ref'],repeated['plan_ref'])
            await asyncio.gather(*list(planner.tasks))
            plan=planner.plans[result['plan_ref']]
            self.assertEqual(plan['status'],'ready')
            self.assertEqual(len(plan['nodes']),1)
            output=await planner.render(result['plan_ref'],[])
            self.assertEqual(output['status'],'complete')
            self.assertEqual('10',Document(output['path']).tables[0].cell(0,1).text)
            with ZipFile(template['_docx']) as old, ZipFile(output['path']) as new:
                self.assertEqual(old.namelist(),new.namelist())
                for name in old.namelist():
                    if name!='word/document.xml': self.assertEqual(old.read(name),new.read(name))
            await services.close()
        asyncio.run(scenario())

    def test_unbound_location_requires_explicit_draft_and_rejects_invented_evidence(self):
        async def scenario():
            template=self.template(missing=True)
            services=RunServices(['first'],env=self.env,catalog=self.catalog)
            services.executor=self.executor
            planner=ReportPlanner(services,template)
            services.report=planner
            result=await planner.prepare({'years':['2025'],'period_mode':'annual','timezone':'Asia/Shanghai'},
                {'sales':{'school':self.school()}})
            await asyncio.gather(*list(planner.tasks))
            data=planner.get(result['plan_ref'],'summary')
            self.assertEqual(data['status'],'ready')
            self.assertEqual(data['coverage']['required_missing'],0)
            self.assertEqual(data['coverage']['drafts_required'],1)
            with self.assertRaisesRegex(DataError,'REPORT_BODY_INCOMPLETE'):
                await planner.render(result['plan_ref'],[])
            self.assertFalse(list((self.executor.context.run_directory/'report').glob('*.docx')))
            with self.assertRaises(DataError):
                await planner.render(result['plan_ref'],[{'section_key':'summary','text':'999','evidence_refs':['invented']}])
            slot=next(item for item in data['slots'] if item.get('draftable'))
            output=await planner.render(result['plan_ref'],[{
                'slot_key':slot['slot_key'],'section_key':'summary',
                'text':'No applicable evidence was retrieved. Confirm the reporting rule.',
                'evidence_refs':[],'evidence_state':'none',
                'gap':'No applicable evidence was retrieved.',
                'next_action':'Confirm the reporting rule.',
            }])
            self.assertEqual(output['status'],'complete_with_data_gaps')
            with self.assertRaises(DataError): planner.get('other-plan')
            await services.close()
        asyncio.run(scenario())

    def test_retired_template_modes_and_unknown_datasets_are_rejected(self):
        template=self.template()
        path=template['_directory']/'template.json'
        original=json.loads(path.read_text())
        for change in ({'missing_policy':'annotated_draft'},{'preserve_structure':False}):
            changed=deepcopy(original); changed['output_policy'].update(change); save(path,changed)
            with self.assertRaisesRegex(DataError,'TEMPLATE_CONTRACT_INVALID'):
                load_template('demo','qa',path.parent.parent)
        save(path,original)
        path=template['_directory']/'document-map.json'
        original=json.loads(path.read_text())
        changed=deepcopy(original)
        changed['locations'][2]['node_type']='table'
        save(path,changed)
        with self.assertRaisesRegex(DataError,'DOCUMENT_MAP_INVALID'):
            load_template('demo','qa',path.parent.parent)
        changed=deepcopy(original)
        changed['locations'][2]['fill']['values']['amount']['dataset_key']='unknown'
        save(path,changed)
        with self.assertRaisesRegex(DataError,'BINDING_INVALID'):
            load_template('demo','qa',path.parent.parent)
        save(path,original)

    def test_explicit_values_reject_private_columns_and_ambiguous_results(self):
        from workflows.writing_docx.report_values import resolve_value
        output=[{'name':'total'},{'name':'internal_id','visibility':'internal_only'}]
        ref=self.executor.results.save([{'total':10,'internal_id':'hidden'}],{'complete':True},output)
        plan={'nodes':[{'dataset_keys':['total'],'status':'succeeded','result_ref':ref}]}
        binding={'kind':'scalar','required':True,'text_template':'${value}',
                 'values':{'value':{'dataset_key':'total','field':'internal_id'}}}
        with self.assertRaisesRegex(DataError,'BINDING_INVALID'):
            resolve_value(binding,plan,self.executor)
        binding['values']['value']['field']='total'
        self.assertEqual(resolve_value(binding,plan,self.executor)['value'],'10')
        self.executor.results.get(ref)['rows'].append({'total':20,'internal_id':'other'})
        with self.assertRaisesRegex(DataError,'BINDING_VALUE_AMBIGUOUS'):
            resolve_value(binding,plan,self.executor)

    def test_template_drift_and_duplicate_locations_fail(self):
        template=self.template()
        path=template['_directory']/'document-map.json'
        document_map=json.loads(path.read_text())
        original=deepcopy(document_map)
        table_location=next(location for location in document_map['locations'] if location['table'])
        self.assertEqual(table_location['table']['table_id'],'T01')
        self.assertEqual(table_location['locator']['path'],
                         '/w:document/w:body/w:tbl/w:tr/w:tc[1]/w:p')
        document_map['locations'].append(document_map['locations'][0])
        save(path,document_map)
        with self.assertRaisesRegex(DataError,'DUPLICATE_SLOT'): load_template('demo','qa',path.parent.parent)
        document_map['locations'].pop()
        table_location['table']['row']=99
        save(path,document_map)
        with self.assertRaisesRegex(DataError,'DOCUMENT_MAP_INVALID'): load_template('demo','qa',path.parent.parent)
        save(path,original)
        document=Document(template['_docx'])
        document.add_paragraph('changed')
        document.save(template['_docx'])
        with self.assertRaisesRegex(DataError,'TEMPLATE_MISMATCH'): load_template('demo','qa',path.parent.parent)

    def test_missing_bound_value_requires_draft_and_other_dataset_gaps_do_not_replace_values(self):
        from workflows.writing_docx.report_values import resolve_value
        binding={'kind':'scalar','required':True,'text_template':'${value}',
                 'evidence_datasets':['total','blocked'],
                 'values':{'value':{'dataset_key':'total','field':'total'}}}
        for rows in ([], [{'total':None}], [{'total':0}]):
            ref=self.executor.results.save(rows,{'complete':True},[{'name':'total'}])
            plan={'nodes':[{'dataset_keys':['total'],'status':'succeeded','result_ref':ref},
                           {'dataset_keys':['blocked'],'status':'blocked','result_ref':None}]}
            result=resolve_value(binding,plan,self.executor)
            if rows and rows[0]['total']==0:
                self.assertEqual(result['value_status'],'filled')
                self.assertEqual(result['value'],'0')
            else:
                self.assertTrue(result['draftable'])
                self.assertNotIn('value',result)
                self.assertEqual(result['missing_values'],['value'])
        plan['nodes'][0].update(status='blocked',result_ref=None)
        self.assertTrue(resolve_value(binding,plan,self.executor)['draftable'])

    def test_supplemental_query_evidence_stays_inside_bound_source_and_scope(self):
        async def scenario():
            template=self.template(missing=True)
            services=RunServices(['first','second'],env=self.env,catalog=self.catalog)
            services.executor=self.executor
            planner=ReportPlanner(services,template)
            services.report=planner
            scope=self.school()
            summary=await planner.prepare({'years':['2025'],'period_mode':'annual','timezone':'Asia/Shanghai'},
                {'sales':{'school':scope}})
            await asyncio.gather(*list(planner.tasks))
            slot=planner.get(summary['plan_ref'],'summary',writing_only=True)['slots'][0]
            extra=self.executor.execute_query_spec('first','sales','total',{},scope)
            self.assertNotIn(extra['result_ref'],slot['evidence_refs'])
            draft={'section_key':'summary','slot_key':slot['slot_key'],'text':'Recorded total: 10.',
                   'evidence_state':'supported','evidence_refs':[extra['result_ref']]}
            self.assertEqual(planner.save_drafts(summary['plan_ref'],[draft])['saved'],1)
            foreign=self.executor.execute_query_spec('second','sales','total',{},self.school(source='second'))
            wrong_scope=self.executor.execute_query_spec('first','sales','total',{},self.school())
            truncated=self.executor.results.save([{'total':10}],
                {**self.executor.results.get(extra['result_ref'])['metadata'],'complete':False},[{'name':'total'}])
            for ref in (foreign['result_ref'],wrong_scope['result_ref'],truncated,'foreign_run_result'):
                with self.subTest(reference=ref), self.assertRaisesRegex(DataError,'EVIDENCE_REQUIRED'):
                    planner.save_drafts(summary['plan_ref'],[{**draft,'evidence_refs':[ref]}])
            await services.close()
        asyncio.run(scenario())

    def test_full_template_requires_body_and_preserves_all_package_parts(self):
        template=load_template('szpt-midterm','document-writing')
        executor=self.make_executor()
        from dataclasses import replace
        executor.context=replace(executor.context,capability_ref='document-writing')
        slots=[{**s,'value_status':'static' if s['kind']=='static' else 'definition_missing',
                **({'draftable':True,'evidence_refs':[]} if s['kind']!='static' else {})}
               for s in template['_slots']]
        with self.assertRaisesRegex(DataError,'REPORT_BODY_INCOMPLETE'):
            render(template,{'plan_ref':'test-draft','slots':slots,'report_parameters':{'years':['2025']}},[],executor)
        drafts=[]
        for slot in slots:
            if slot['kind']=='static':
                continue
            draft={'section_key':slot['section_key'],'slot_key':slot['slot_key'],
                   'text':'Evidence is missing.','evidence_refs':[],'evidence_state':'none',
                   'gap':'Evidence is missing.'}
            if slot['kind']=='narrative':
                draft.update(text='Evaluate the original subject. Evidence is missing. Collect dated records.',
                             analysis_basis='Evaluate the original subject.',
                             next_action='Collect dated records.')
            drafts.append(draft)
        output=render(template,{'plan_ref':'test-draft','slots':slots,'report_parameters':{'years':['2025']}},drafts,executor)
        self.assertEqual(output['status'],'complete_with_data_gaps')
        before=document_xml(template['_docx'])
        after=document_xml(output['path'])
        self.assertTrue(after.xpath('.//w:color[@w:val="FFC000"]',namespaces=NS))
        for tag in ('tbl','tr','tc','p','sectPr','drawing'):
            self.assertEqual(len(before.xpath(f'.//w:{tag}',namespaces=NS)),len(after.xpath(f'.//w:{tag}',namespaces=NS)))
        with ZipFile(template['_docx']) as old, ZipFile(output['path']) as new:
            for name in old.namelist():
                if name!='word/document.xml': self.assertEqual(old.read(name),new.read(name))

    def test_mixed_runs_preserve_markup_and_replace_text(self):
        from lxml import etree
        from workflows.writing_docx.report_renderer import replace_paragraph
        markup='<w:p xmlns:w="'+NS['w']+'"><w:pPr/><w:r><w:rPr><w:b/></w:rPr><w:t>abc</w:t></w:r><w:bookmarkStart w:id="1" w:name="x"/><w:r><w:t>def</w:t></w:r><w:bookmarkEnd w:id="1"/></w:p>'
        for replacement in ('abc123def','ABCdef','abcdefXYZ','','hello world','aXYcdeZ'):
            node=etree.fromstring(markup)
            replace_paragraph(node,replacement)
            self.assertEqual(''.join(node.xpath('.//w:t/text()',namespaces=NS)),replacement)
            self.assertEqual(len(node.xpath('.//w:r',namespaces=NS)),2)
            self.assertEqual(len(node.xpath('.//w:b|.//w:bookmarkStart|.//w:bookmarkEnd',namespaces=NS)),3)

    def test_unregistered_alternate_template_rejected(self):
        with self.assertRaises(DataError): load_template('double-high-annual','document-writing')

    def test_fill_markers_only_clear_registered_editing_colors(self):
        from lxml import etree
        from workflows.writing_docx.report_renderer import clear_fill_markers
        markup = '<w:p xmlns:w="' + NS['w'] + '"><w:r><w:rPr><w:b/><w:sz w:val="24"/><w:highlight w:val="yellow"/><w:color w:val="ee0000"/></w:rPr><w:t>Filled</w:t></w:r><w:r><w:rPr><w:color w:val="FF0000"/></w:rPr><w:t>Value</w:t></w:r><w:r><w:rPr><w:color w:val="0070C0"/></w:rPr><w:t>Brand</w:t></w:r></w:p>'
        node = etree.fromstring(markup)
        clear_fill_markers(node)
        self.assertEqual(node.xpath('.//w:color/@w:val', namespaces=NS), ['0070C0'])
        self.assertFalse(node.xpath('.//w:highlight', namespaces=NS))
        self.assertEqual(len(node.xpath('.//w:b|.//w:sz|.//w:r|.//w:t', namespaces=NS)), 8)
        self.assertEqual(''.join(node.xpath('.//w:t/text()', namespaces=NS)), 'FilledValueBrand')

    def test_fidelity_rejects_changed_run_properties(self):
        template=self.template()
        document=Document(template['_docx'])
        document.paragraphs[0].runs[0].bold=True
        changed=self.root/'changed.docx'
        document.save(changed)
        with self.assertRaisesRegex(DataError,'TEMPLATE_STRUCTURE_CHANGED'):
            check_fidelity(template,changed)

    def test_fidelity_allows_marker_cleanup_but_not_other_format_changes(self):
        from lxml import etree
        from workflows.writing_docx.report_renderer import clear_fill_markers
        template = self.template()
        template['clear_fill_markers'] = True
        xml = document_xml(template['_docx'])
        slot = next(s for s in template['_slots'] if s['kind'] != 'static')
        paragraph = xml.xpath(slot['locator']['path'], namespaces=NS)[0]
        properties = etree.SubElement(paragraph.find('w:r', NS), '{' + NS['w'] + '}rPr')
        color = etree.SubElement(properties, '{' + NS['w'] + '}color')
        color.set('{' + NS['w'] + '}val', 'EE0000')
        with ZipFile(template['_docx']) as archive:
            members = [(m, archive.read(m.filename)) for m in archive.infolist()]
        def write(path):
            with ZipFile(path, 'w') as archive:
                for member, data in members:
                    archive.writestr(member, etree.tostring(xml) if member.filename == 'word/document.xml' else data)
        write(template['_docx'])
        clear_fill_markers(paragraph)
        changed = self.root / 'cleaned.docx'
        write(changed)
        self.assertTrue(check_fidelity(template, changed)['fill_markers_cleared'])
        etree.SubElement(properties, '{' + NS['w'] + '}b')
        write(changed)
        with self.assertRaisesRegex(DataError, 'TEMPLATE_STRUCTURE_CHANGED'):
            check_fidelity(template, changed)

    def test_file_name_follows_report_parameters_and_rejects_paths(self):
        from workflows.writing_docx.report_renderer import report_file_name
        template = {'file_name': 'Report-${years}.docx'}
        self.assertEqual(report_file_name(template, {'years': ['2026']}), 'Report-2026.docx')
        for parameters in ({}, {'years': ['../2025']}, {'years': ['a\\b']}):
            with self.assertRaisesRegex(DataError, 'REPORT_FILE_NAME_INVALID'):
                report_file_name(template, parameters)

    def test_visual_review_requires_reading_every_page_and_current_file(self):
        from types import SimpleNamespace
        template=self.template()
        rendered={'path':str(template['_docx'])}
        validation={'validation_ref':'v','document_sha256':sha256(template['_docx'].read_bytes()).hexdigest(),
                    'page_count':2,'pages_read':[1]}
        planner=SimpleNamespace(latest='p',plans={'p':{'rendered':rendered,'validation':validation}})
        with self.assertRaisesRegex(DataError,'REPORT_VALIDATION_REQUIRED'):
            validated_plan(planner,'p','other')
        validated_plan(planner,'p','v')
        with self.assertRaisesRegex(DataError,'REPORT_PAGES_NOT_READ'):
            record_page_review(validation,[2],True,'checked')
        self.assertEqual(record_page_review(validation,[1],True,'checked')['visual_review'],'review_required')
        validation['pages_read']=[1,2]
        self.assertEqual(record_page_review(validation,[2],False,'clipping')['visual_review'],'review_required')
        self.assertEqual(record_page_review(validation,[2],True,'fixed')['visual_review'],'passed')
        template['_docx'].write_bytes(b'changed')
        with self.assertRaisesRegex(DataError,'REPORT_VALIDATION_REQUIRED'):
            validated_plan(planner,'p','v')

    def test_saved_body_batches_invalidate_render_and_validation(self):
        async def scenario():
            template=self.template()
            binding=template['_slots'][-1]
            binding.update(kind='narrative',required=True,evidence_datasets=['total'])
            services=RunServices(['first'],env=self.env,catalog=self.catalog)
            services.executor=self.executor
            planner=ReportPlanner(services,template)
            services.report=planner
            result=await planner.prepare({'years':['2025'],'period_mode':'annual','timezone':'Asia/Shanghai'},
                {'sales':{'school':self.school()}})
            await asyncio.gather(*list(planner.tasks))
            reference=result['plan_ref']
            data=planner.get(reference,'summary',writing_only=True)
            self.assertEqual(len(data['slots']),1)
            slot=data['slots'][0]
            draft={'slot_key':slot['slot_key'],'section_key':'summary','text':'The recorded total is 10.',
                   'evidence_refs':slot['evidence_refs'],'evidence_state':'supported'}
            self.assertEqual(planner.save_drafts(reference,[draft])['saved'],1)
            plan=planner.plans[reference]
            plan.update(rendered={'path':'old'},validation={'validation_ref':'old'})
            planner.save_drafts(reference,[{**draft,'text':'The total remains 10.'}])
            self.assertNotIn('rendered',plan)
            self.assertNotIn('validation',plan)
            with self.assertRaisesRegex(DataError,'EVIDENCE_REQUIRED'):
                planner.save_drafts(reference,[{**draft,'evidence_refs':['foreign_result']}])
            with self.assertRaisesRegex(DataError,'DRAFT_NUMBER_UNSUPPORTED'):
                planner.save_drafts(reference,[{**draft,'text':'Recorded total 9999.'}])
            with self.assertRaisesRegex(DataError,'DRAFT_DUPLICATE'):
                planner.save_drafts(reference,[draft,draft])
            self.assertEqual(len(plan['drafts']),1)
            await services.close()
        asyncio.run(scenario())

    def test_bind_enforces_selected_sources_and_frozen_template_revision(self):
        async def scenario():
            template=self.template()
            self.config['policies']['first']['templates']=['demo']
            self.save_connections()
            payload={'run_id':'selected-template', 'capability_ref':'qa', '_template_key':'demo',
                '_data_identity':{'tenant_id':'jwt-a','user_id':'u1'},
                '_data_run_directory':str(self.root/'selected-template'),
                '_workflow_assets':{'template_revision':template['_revision']}}
            services=RunServices(['first','second'],env=self.env,catalog=self.catalog)
            with patch('workflows.writing_docx.template_assets.load_template',return_value=template):
                await services.bind(payload)
                self.assertEqual([s['source_key'] for s in services.current().list_data_sources()['sources']],['first'])
                with self.assertRaises(DataError): services.current().access('second')
                await services.close()
                with self.assertRaisesRegex(DataError,'TEMPLATE_MISMATCH'):
                    await services.bind({**payload,'_workflow_assets':{'template_revision':'old-revision'}})
                self.assertIsNone(services.executor)
        asyncio.run(scenario())

    def test_evidence_modes_require_visible_gaps_and_reject_unfounded_numbers(self):
        from workflows.writing_docx.report_renderer import validate_drafts
        ref=self.executor.results.save([{'total':10}],{'complete':True},[{'name':'total','type':'number'}])
        empty=self.executor.results.save([],{'complete':True},[{'name':'total','type':'number'}])
        plan={'slots':[{'slot_key':'body','section_key':'summary','kind':'narrative','required':True,
                        'draftable':True,'evidence_refs':[ref,empty]}],
              'report_parameters':{}}
        draft={'slot_key':'body','section_key':'summary','text':'Analyze coverage. Records missing. Collect evidence.',
               'evidence_refs':[],'evidence_state':'none','analysis_basis':'Analyze coverage.',
               'gap':'Records missing.','next_action':'Collect evidence.'}
        self.assertIn('body',validate_drafts(plan,[draft],self.executor))
        with self.assertRaisesRegex(DataError,'EVIDENCE_GAP_REQUIRED'):
            validate_drafts(plan,[{**draft,'text':'Records missing.'}],self.executor)
        with self.assertRaisesRegex(DataError,'EVIDENCE_REQUIRED'):
            validate_drafts(plan,[{**draft,'evidence_state':'supported','evidence_refs':[empty]}],self.executor)
        limited={**draft,'text':'Recorded 10. Historical records missing.', 'evidence_state':'limited',
                 'gap':'Historical records missing.','evidence_refs':[ref]}
        self.assertIn('body',validate_drafts(plan,[limited],self.executor))
        with self.assertRaisesRegex(DataError,'EVIDENCE_STATE_CONFLICT'):
            validate_drafts(plan,[{**limited,'evidence_state':'supported'}],self.executor)
        with self.assertRaisesRegex(DataError,'DRAFT_NUMBER_UNSUPPORTED'):
            validate_drafts(plan,[{**limited,'text':'Recorded 999. Historical records missing.'}],self.executor)

    def test_yellow_color_preserves_style_and_disappears_when_evidence_is_supported(self):
        from workflows.writing_docx.report_renderer import mark_uncertainty
        from lxml import etree
        template=self.template()
        xml=document_xml(template['_docx'])
        slot=next(s for s in template['_slots'] if s['kind']!='static')
        paragraph=xml.xpath(slot['locator']['path'],namespaces=NS)[0]
        mark_uncertainty(paragraph)
        changed=self.root/'yellow.docx'
        with ZipFile(template['_docx']) as source,ZipFile(changed,'w') as target:
            for member in source.infolist():
                target.writestr(member,etree.tostring(xml) if member.filename=='word/document.xml' else source.read(member.filename))
        self.assertTrue(check_fidelity(template,changed)['unchanged_layout_and_run_properties'])
        # Rendering always starts from the original, so confirmed values cannot inherit old yellow.
        result=self.executor.execute_query_spec('first','sales','total',{},self.school())
        plan={'plan_ref':'supported','slots':[{'slot_key':slot['slot_key'],'value_status':'filled','value':'10'}]}
        output=render(template,plan,[],self.executor)
        self.assertFalse(document_xml(output['path']).xpath('.//w:color[@w:val="FFC000"]',namespaces=NS))

    def test_report_mcp_publish_gate_and_page_images(self):
        async def scenario():
            from types import SimpleNamespace
            from workflows.writing_docx.tools import create_reports_server
            from base64 import b64decode
            template=self.template()
            services=RunServices(['first'],env=self.env,catalog=self.catalog)
            services.executor=self.executor
            planner=ReportPlanner(services,template)
            services.report=planner
            summary=await planner.prepare({'years':['2025'],'period_mode':'annual','timezone':'Asia/Shanghai'},
                {'sales':{'school':self.school()}})
            await asyncio.gather(*list(planner.tasks))
            ref=summary['plan_ref']
            output=await planner.render(ref,[])
            services.artifact_directories={'deliverables_directory':str(self.root/'deliverables'),'session_directory':str(self.root/'session')}
            page=self.root/'page.png'
            page.write_bytes(b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+aX1sAAAAASUVORK5CYII='))
            plan=planner.plans[ref]
            plan['validation']={'validation_ref':'v','document_sha256':sha256(Path(output['path']).read_bytes()).hexdigest(),
                                'page_count':2,'pages':[str(page),str(page)],'visual_review':'rendered_pages_require_review'}
            entry=create_reports_server(services)['instance'].get_request_handler('tools/call')
            async def call(name,**args):
                result=await entry.handler(None,entry.params_type(name=name,arguments=args))
                return result
            args={'plan_ref':ref,'validation_ref':'v'}
            failed=await call('publish_report',**args)
            self.assertIn('REPORT_VISUAL_REVIEW_REQUIRED',failed.content[0].text)
            self.assertFalse((self.root/'deliverables/report.docx').exists())
            images=await call('read_report_pages',**args,page_numbers=[1,2])
            self.assertEqual(sum(c.type=='image' for c in images.content),2)
            await call('review_report_pages',**args,page_numbers=[1],passed=True,notes='first page checked')
            self.assertIn('REPORT_VISUAL_REVIEW_REQUIRED',(await call('publish_report',**args)).content[0].text)
            await call('review_report_pages',**args,page_numbers=[2],passed=True,notes='second page checked')
            result=json.loads((await call('publish_report',**args)).content[0].text)
            self.assertTrue(result['published'])
            self.assertEqual(result['layout_validation'],'structure_and_all_pages_reviewed')
            self.assertTrue((self.root/'deliverables/report.docx').exists())
            await services.close()
        asyncio.run(scenario())


if __name__ == '__main__': unittest.main()
