"""Template selection, frozen instructions and client reuse must agree."""

from copy import deepcopy
from hashlib import sha256
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from docx import Document

from data_access.context import DataError
from runtime.config import build_options, prepare_workflow_assets, workflow_prompt_documents
from runtime.data_services import RunServices
from runtime.prompt_documents import MAX_DOCUMENT_BYTES, read_documents
from runtime.protocol import AgentRunRequest
from workflows.writing_docx.template_assets import load_template, stage_template


class TemplateDocumentTests(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        self.root = Path(temp.name)
        self.directory = self.root / 'workflow'
        self.directory.mkdir()
        self.workflow = {'name':'writing-docx', '_directory':str(self.directory),
                         '_body':'COMMON_RULES',
                         'templates':{}, 'data_access':'optional'}
        for name in ('alpha','beta'):
            directory = self.directory / 'templates' / name
            directory.mkdir(parents=True)
            path = directory / 'template.docx'
            Document().save(path)
            self.save(directory/'template.json', {'template_key':name, 'enabled':True,
                'capabilities':['document-writing'], 'version':1,
                'assets':{'docx':{'file':'template.docx','sha256':sha256(path.read_bytes()).hexdigest()},
                    'writing_guide':'writing-guide.md'},
                'data':{'source_roles':{'hpm':'source-a'},'scope_roles':{}},
                'report':{'file_name':'report.docx'}})
            (directory/'writing-guide.md').write_text('REPORT_ADAPTER\n'+name.upper()+'_GUIDE',encoding='utf-8')
            self.workflow['templates'][name] = f'templates/{name}/template.json'
        for name, value in [('runtime.config.PROJECT_ROOT', self.root),
                            ('runtime.config.load_workflow_config', None)]:
            patcher = patch(name, return_value=self.workflow) if value is None else patch(name,value)
            patcher.start(); self.addCleanup(patcher.stop)
        patcher = patch('workflows.writing_docx.template_assets.load_template', side_effect=self.load)
        patcher.start(); self.addCleanup(patcher.stop)

    @staticmethod
    def save(path, data):
        path.write_text(json.dumps(data),encoding='utf-8')

    def load(self, name, capability):
        return load_template(name,capability,self.directory/'templates')

    def payload(self, name=None):
        payload={'workflow_name':'writing-docx','capability_ref':'document-writing'}
        if name is not None: payload['_template_key']=name
        return payload

    def test_selected_guide_only_and_no_template_for_uploaded_reference(self):
        for name, present, absent in [('alpha','ALPHA_GUIDE','BETA_GUIDE'),('beta','BETA_GUIDE','ALPHA_GUIDE')]:
            prompt=prepare_workflow_assets(self.payload(name))['prompt']
            self.assertIn('COMMON_RULES',prompt)
            self.assertIn('REPORT_ADAPTER',prompt)
            self.assertIn(present,prompt)
            self.assertNotIn(absent,prompt)
        with patch('runtime.config.Catalog') as catalog, patch.dict('os.environ',{},clear=True):
            catalog.return_value.sources_for.return_value=[]
            options=build_options(self.payload())
        self.assertIn('COMMON_RULES',options.system_prompt['append'])
        for marker in ('ALPHA_GUIDE','BETA_GUIDE','REPORT_ADAPTER'):
            self.assertNotIn(marker,options.system_prompt['append'])
        self.assertNotIn('reports',options.mcp_servers)
        self.assertEqual(options.mcp_servers['office']['command'],'officecli')
        self.assertEqual(options.mcp_servers['office']['args'],['mcp'])
        self.assertIn('documents',options.mcp_servers)
        self.assertNotIn('docx',options.mcp_servers)

    def test_declared_documents_order_and_unregistered_file_exclusion(self):
        directory=self.directory/'templates/alpha'
        config=json.loads((directory/'template.json').read_text())
        config['assets']['writing_guide']='second.md'
        self.save(directory/'template.json',config)
        (directory/'second.md').write_text('SECOND_GUIDE')
        (directory/'unregistered.md').write_text('HIDDEN_GUIDE')
        prompt=prepare_workflow_assets(self.payload('alpha'))['prompt']
        self.assertIn('SECOND_GUIDE',prompt)
        self.assertNotIn('ALPHA_GUIDE',prompt)
        self.assertNotIn('HIDDEN_GUIDE',prompt)

    def test_guide_mutation_changes_next_run_revision_but_not_frozen_prompt(self):
        payload=self.payload('alpha')
        before=deepcopy(prepare_workflow_assets(payload))
        (self.directory/'templates/alpha/writing-guide.md').write_text('UPDATED_GUIDE')
        self.assertEqual(prepare_workflow_assets(payload),before)
        after=prepare_workflow_assets(self.payload('alpha'))
        self.assertNotEqual(before['revision'],after['revision'])
        self.assertNotEqual(before['template_revision'],after['template_revision'])
        self.assertIn('UPDATED_GUIDE',after['prompt'])
        # Reusing the same business payload with a new selection cannot reuse the old guide.
        payload['_template_key']='beta'
        self.assertIn('BETA_GUIDE',prepare_workflow_assets(payload)['prompt'])

    def test_common_rule_mutation_changes_client_fingerprint_and_payload_stays_serializable(self):
        from server import _client_config_fingerprint
        request=AgentRunRequest.from_dict({'protocol':'agent-run/v1','runId':'run-test',
            'messageId':'msg-test','businessSessionId':'session-test','capabilityRef':'document-writing',
            'input':{'text':'write'}})
        with patch('server._data_config_fingerprint',return_value=[]):
            old=self.payload('alpha')
            before=_client_config_fingerprint(request,old)
            json.dumps(old)
            self.workflow['_body'] = 'CHANGED_COMMON'
            self.assertEqual(before,_client_config_fingerprint(request,old))
            self.assertNotEqual(before,_client_config_fingerprint(request,self.payload('alpha')))
            self.assertNotEqual(before,_client_config_fingerprint(request,self.payload('beta')))

    def test_invalid_guides_are_rejected_before_agent_start(self):
        directory=self.directory/'templates/alpha'
        original=json.loads((directory/'template.json').read_text())
        invalid=[None,'missing.md','../beta/writing-guide.md',
                 'private.not_for_model.md',str(directory/'writing-guide.md'),'config.json']
        for names in invalid:
            with self.subTest(names=names):
                changed=deepcopy(original); changed['assets']['writing_guide']=names
                self.save(directory/'template.json',changed)
                with self.assertRaisesRegex(DataError,'TEMPLATE_DOCUMENT_INVALID'):
                    prepare_workflow_assets(self.payload('alpha'))
        self.save(directory/'template.json',original)
        for content in ('', 'a'*(MAX_DOCUMENT_BYTES+1)):
            (directory/'writing-guide.md').write_text(content)
            with self.assertRaisesRegex(DataError,'TEMPLATE_DOCUMENT_INVALID'):
                self.load('alpha','document-writing')

    def test_rejects_cross_platform_absolute_paths_and_total_prompt_overflow(self):
        for name in ('C:/secret.md', '//server/share/secret.md'):
            with self.assertRaises(RuntimeError): read_documents(self.directory,[name])
        template=self.load('alpha','document-writing')
        template['_documents']=[{'name':str(n),'text':'a'*23000} for n in range(4)]
        with self.assertRaisesRegex(RuntimeError,'总量'):
            workflow_prompt_documents(self.workflow,template=template)

    def test_disabled_or_wrong_capability_template_is_rejected(self):
        with self.assertRaisesRegex(DataError,'TEMPLATE_FORBIDDEN'):
            self.load('alpha','conversation')
        path=self.directory/'templates/alpha/template.json'
        config=json.loads(path.read_text()); config['enabled']=False; self.save(path,config)
        with self.assertRaisesRegex(DataError,'TEMPLATE_FORBIDDEN'):
            prepare_workflow_assets(self.payload('alpha'))

    def test_staged_copy_is_repeatable_and_never_overwrites_edits(self):
        template=self.load('alpha','document-writing')
        source=template['_docx'].read_bytes()
        staged=stage_template(template,self.root/'run-a')
        self.assertEqual(staged.read_bytes(),source)
        self.assertEqual(stage_template(template,self.root/'run-a'),staged)
        other=stage_template(template,self.root/'run-b')
        self.assertNotEqual(staged,other)
        staged.write_bytes(b'changed')
        with self.assertRaisesRegex(DataError,'TEMPLATE_COPY_MODIFIED'):
            stage_template(template,self.root/'run-a')
        self.assertEqual(staged.read_bytes(),b'changed')
        self.assertEqual(template['_docx'].read_bytes(),source)
        self.assertEqual(other.read_bytes(),source)

    def test_changed_source_cannot_be_staged_with_frozen_digest(self):
        template=self.load('alpha','document-writing')
        template['_docx'].write_bytes(b'changed')
        with self.assertRaisesRegex(DataError,'TEMPLATE_MISMATCH'):
            stage_template(template,self.root/'run')

    def test_registered_template_requires_run_workspace(self):
        with patch('runtime.config.create_run_services',return_value=RunServices([])):
            with self.assertRaisesRegex(DataError,'TEMPLATE_WORKSPACE_REQUIRED'):
                build_options(self.payload('alpha'))

    def test_draft_publication_preserves_source_and_staged_reference(self):
        from tools.artifacts import publish_artifact
        source=self.directory/'templates/alpha/template.docx'
        document=Document()
        document.add_heading('Construction progress',1)
        document.add_paragraph('Template example')
        document.add_table(rows=1,cols=2).cell(0,0).text='Metric'
        document.save(source)
        original=source.read_bytes()
        work=self.root/'session/work'
        staged=stage_template(self.load('alpha','document-writing'),work)
        self.assertEqual(Document(staged).paragraphs[0].text,'Construction progress')
        self.assertEqual(len(Document(staged).tables),1)
        draft=Document(staged)
        draft.paragraphs[1].text='Evidence still required'
        from docx.shared import RGBColor
        draft.paragraphs[1].runs[0].font.color.rgb=RGBColor(255,192,0)
        output=work/'draft.docx'
        draft.save(output)
        result=publish_artifact(output,'report.docx',work,self.root/'session/deliverables',self.root/'session')
        self.assertEqual(source.read_bytes(),original)
        self.assertEqual(staged.read_bytes(),original)
        self.assertTrue(result)
        published=self.root/'session/deliverables'/result['artifactId']/'content'
        self.assertEqual(Document(published).paragraphs[1].text,'Evidence still required')
        self.assertEqual(str(Document(published).paragraphs[1].runs[0].font.color.rgb),'FFC000')

    def test_source_change_after_prompt_freeze_rejects_agent_start(self):
        payload=self.payload('alpha')
        prepare_workflow_assets(payload)
        payload.update(session_directory=str(self.root/'session'),work_directory=str(self.root/'session/work'),
                       deliverables_directory=str(self.root/'session/output'))
        (self.directory/'templates/alpha/writing-guide.md').write_text('NEW_RULES')
        with patch('runtime.config.create_run_services',return_value=RunServices([])):
            with self.assertRaisesRegex(DataError,'TEMPLATE_MISMATCH'):
                build_options(payload)
