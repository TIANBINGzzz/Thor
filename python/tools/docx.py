"""Small, session-scoped DOCX inspection and template tools for the Agent SDK."""

from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from runtime.claude_sdk import create_sdk_mcp_server, sdk_tool


def _roots(base_dir: str | Path, additional_dirs: list[str] | None) -> list[Path]:
    values = [Path(base_dir), *(Path(item) for item in (additional_dirs or []))]
    return [item.expanduser().resolve() for item in values]


def _inside(path: Path, roots: list[Path]) -> bool:
    return any(path == root or root in path.parents for root in roots)


def _resolve(path_value: str, roots: list[Path], *, must_exist: bool = True) -> Path:
    """按允许目录解析输入 DOCX 路径并返回绝对路径，越界、扩展名错误或缺失时抛出异常。"""
    if not isinstance(path_value, str) or not path_value.strip():
        raise ValueError("path 不能为空")
    path = Path(path_value).expanduser()
    if not path.is_absolute():
        path = roots[0] / path
    path = path.resolve()
    if not _inside(path, roots):
        raise ValueError("DOCX 路径必须位于当前项目或会话目录内")
    if must_exist and not path.is_file():
        raise FileNotFoundError(f"文件不存在：{path.name}")
    if path.suffix.lower() != ".docx":
        raise ValueError("只支持 .docx 文件")
    return path


def _body_paragraphs(parent):
    for block in parent.iter_inner_content():
        if isinstance(block, Paragraph):
            yield block
        elif isinstance(block, Table):
            seen = set()
            for row in block.rows:
                for cell in row.cells:
                    if cell._tc not in seen:
                        seen.add(cell._tc)
                        yield from _body_paragraphs(cell)


def _paragraphs(document: Document):
    """Read in document order, including nested tables without repeating merged cells."""
    yield from _body_paragraphs(document)
    for section in document.sections:
        yield from section.header.paragraphs
        yield from section.footer.paragraphs


def _semantic_blocks(parent):
    blocks = []
    for block in parent.iter_inner_content():
        if isinstance(block, Paragraph):
            style = block.style
            level = block._p.xpath('./w:pPr/w:outlineLvl/@w:val')
            visited = set()
            inherited = style
            while not level and inherited is not None and inherited.style_id not in visited:
                visited.add(inherited.style_id)
                level = inherited.element.xpath('./w:pPr/w:outlineLvl/@w:val')
                inherited = inherited.base_style
            blocks.append({'type':'paragraph', 'text':block.text, 'style':style.name if style else None,
                           'heading_level':int(level[0]) + 1 if level and int(level[0]) < 9 else None})
        else:
            seen, rows = set(), []
            for row in block.rows:
                cells = []
                for column, cell in enumerate(row.cells, 1):
                    if cell._tc in seen:
                        continue
                    seen.add(cell._tc)
                    span = cell._tc.xpath('./w:tcPr/w:gridSpan/@w:val')
                    cells.append({'column':column, 'column_span':int(span[0]) if span else 1,
                                  'blocks':_semantic_blocks(cell)})
                rows.append(cells)
            blocks.append({'type':'table', 'style':block.style.name if block.style else None, 'rows':rows})
    return blocks


def _document_summary(path: Path, document: Document) -> dict[str, Any]:
    """接收文件路径和 Word 文档，返回段落、表格、计数及模板占位符的结构化摘要。"""
    paragraphs = [p.text for p in _paragraphs(document)]
    tables = [
        [[cell.text for cell in row.cells] for row in table.rows]
        for table in document.tables
    ]
    placeholders = sorted({
        match
        for text in paragraphs
        for match in re.findall(r"\{\{\s*([^{}]+?)\s*\}\}", text)
    })
    return {
        "path": str(path),
        "paragraphs": paragraphs,
        "paragraphCount": len(paragraphs),
        "tables": tables,
        "tableCount": len(tables),
        "placeholders": placeholders,
        "structure": _semantic_blocks(document),
        "sections": [{'width':s.page_width, 'height':s.page_height,
                      'header':[p.text for p in s.header.paragraphs],
                      'footer':[p.text for p in s.footer.paragraphs]} for s in document.sections],
    }


def _text_result(value: Any) -> dict[str, list[dict[str, str]]]:
    return {"content": [{"type": "text", "text": json.dumps(value, ensure_ascii=False)}]}


def inspect_document(path_value: str, base_dir: str | Path, additional_dirs: list[str] | None = None) -> dict[str, Any]:
    """在指定允许目录内读取 DOCX，返回段落、表格及占位符摘要字典。"""
    roots = _roots(base_dir, additional_dirs)
    path = _resolve(path_value, roots)
    return _document_summary(path, Document(path))


def extract_document(path_value: str, base_dir: str | Path, additional_dirs: list[str] | None = None) -> dict[str, Any]:
    """按文件路径和允许目录提取 DOCX 段落文本，返回包含路径及合并纯文本的字典。"""
    roots = _roots(base_dir, additional_dirs)
    path = _resolve(path_value, roots)
    return {"path": str(path), "text": "\n".join(p.text for p in _paragraphs(Document(path)))}


def replace_document(
    path_value: str,
    replacements: dict[str, Any],
    base_dir: str | Path,
    additional_dirs: list[str] | None = None,
    output_path: str | None = None,
) -> dict[str, Any]:
    """接收模板路径、占位符替换映射和可选输出路径，替换内容并另存为新 DOCX。

    返回输出路径、发生修改的段落数和替换键列表，输出路径必须区别于源模板。
    """
    roots = _roots(base_dir, additional_dirs)
    source = _resolve(path_value, roots)
    if not isinstance(replacements, dict) or not replacements:
        raise ValueError("replacements 必须是非空对象")
    output = _resolve(output_path, roots, must_exist=False) if output_path else source.with_name(f"{source.stem}.generated.docx")
    if output == source:
        raise ValueError("输出文件必须不同于模板文件")
    output.parent.mkdir(parents=True, exist_ok=True)
    document = Document(source)
    changed = 0
    tokens = {
        (key if str(key).startswith("{{") and str(key).endswith("}}") else "{{" + str(key).strip("{} ") + "}}"): str(value)
        for key, value in replacements.items()
    }
    for paragraph in _paragraphs(document):
        original = paragraph.text
        updated = original
        for token, value in tokens.items():
            updated = updated.replace(token, value)
        if updated != original:
            # Keep run styling when a placeholder is contained in one run;
            # use paragraph text only as a fallback for split-run templates.
            if any(token in run.text for token in tokens for run in paragraph.runs):
                for run in paragraph.runs:
                    for token, value in tokens.items():
                        run.text = run.text.replace(token, value)
            else:
                paragraph.text = updated
            changed += 1
    document.save(output)
    return {"path": str(output), "changedParagraphs": changed, "replacements": list(replacements)}


def create_docx_server(base_dir: str | Path, additional_dirs: list[str] | None = None):
    """接收基础目录和额外允许目录，返回注册检查、提取和替换工具的进程内 MCP 服务配置。"""
    roots = _roots(base_dir, additional_dirs)

    @sdk_tool(
        "docx_inspect",
        "按文档顺序读取DOCX标题层级、样式、嵌套表格和页眉页脚；返回语义结构供参考撰写，不生成位置地图。",
        {"type": "object", "properties": {"path": {"type": "string"}}, "required": ["path"]},
    )
    async def inspect_docx(args: dict[str, Any]) -> dict[str, Any]:
        """接收含 path 的工具参数，返回以 MCP 文本格式封装的 DOCX 结构摘要。"""
        return _text_result(inspect_document(args.get("path", ""), roots[0], [str(root) for root in roots[1:]]))

    @sdk_tool(
        "docx_extract_text",
        "提取 DOCX 正文、表格、页眉和页脚的纯文本。",
        {"type": "object", "properties": {"path": {"type": "string"}}, "required": ["path"]},
    )
    async def extract_docx(args: dict[str, Any]) -> dict[str, Any]:
        """接收含 path 的工具参数，返回以 MCP 文本格式封装的文档纯文本。"""
        return _text_result(extract_document(args.get("path", ""), roots[0], [str(root) for root in roots[1:]]))

    @sdk_tool(
        "docx_replace_text",
        "按映射替换 DOCX 中的 {{placeholder}}，并保存为新文件。",
        {
            "type": "object",
            "properties": {
                "path": {"type": "string"},
                "replacements": {"type": "object", "additionalProperties": {"type": "string"}},
                "output_path": {"type": "string"},
            },
            "required": ["path", "replacements"],
        },
    )
    async def replace_docx(args: dict[str, Any]) -> dict[str, Any]:
        """接收路径、替换映射和可选输出路径，返回以 MCP 文本格式封装的新文件生成结果。"""
        return _text_result(replace_document(
            args.get("path", ""), args.get("replacements"), roots[0],
            [str(root) for root in roots[1:]], args.get("output_path"),
        ))

    return create_sdk_mcp_server("docx", version="0.1.0", tools=[inspect_docx, extract_docx, replace_docx])
